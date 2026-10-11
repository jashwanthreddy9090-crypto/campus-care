from flask import Flask, render_template_string, request, redirect, url_for, session, flash, send_from_directory, jsonify, Response, abort
import sqlite3, os, uuid, re, json, logging, base64, io, csv
from urllib.parse import urlparse, quote_plus
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename

app = Flask(__name__)
application = app  # WSGI alias for deployment platforms
app.secret_key = os.environ.get("SECRET_KEY", "")
if not app.secret_key:
    # A stable secret is mandatory in production; local development gets a warning-safe generated fallback.
    app.secret_key = os.urandom(32).hex()
    if os.environ.get("VERCEL"):
        raise RuntimeError("Set SECRET_KEY in Vercel Environment Variables before deploying.")
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax", SESSION_COOKIE_SECURE=bool(os.environ.get("VERCEL")))
logging.basicConfig(level=logging.INFO)
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATABASE_URL = (os.environ.get("DATABASE_URL") or os.environ.get("POSTGRES_URL") or os.environ.get("POSTGRES_PRISMA_URL") or "").strip()
DATABASE = os.path.join("/tmp", "campuscare.db") if os.environ.get("VERCEL") else os.path.join(BASE_DIR, "campuscare.db")
UPLOAD_FOLDER = os.path.join("/tmp", "campuscare_uploads") if os.environ.get("VERCEL") else os.path.join(BASE_DIR, "static", "uploads")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config["MAX_CONTENT_LENGTH"] = 6 * 1024 * 1024  # 6 MB total request limit

CATEGORIES = ["Hostel","Classroom","Laboratory","Washroom","Electricity","Water","Wi-Fi","Food/Mess","Cleaning","Security","Other"]
PRIORITIES = ["Low","Medium","High","Emergency"]
STATUSES = ["Submitted","Assigned","In Progress","Resolved","Closed"]
ALLOWED_EXTENSIONS = {"png","jpg","jpeg","gif","webp"}

def valid_campus_email(email):
    return bool(re.fullmatch(r"[^@\s]+@campus\.360", email.strip().lower()))

DEMO_WORKERS = [
    ("Ravi Kumar", "Electrical", "9876543210"),
    ("Suresh Reddy", "Plumbing", "9876543211"),
    ("Arjun Kumar", "Cleaning", "9876543212"),
    ("Naveen Kumar", "Maintenance", "9876543213"),
    ("Mahesh", "Wi-Fi & Network", "9876543214"),
    ("Prakash", "Security", "9876543215"),
]

CONTACTS = [
    ("Campus Admin", "Administration", "9876500001"),
    ("Hostel Warden", "Hostel Office", "9876500002"),
    ("Maintenance Desk", "Maintenance", "9876500003"),
    ("Security Desk", "Security", "9876500004"),
    ("Emergency Help", "Emergency", "112"),
]

class PostgresConnection:
    """Small SQLite-compatible adapter for managed PostgreSQL (Neon/Vercel Postgres)."""
    def __init__(self, dsn):
        try:
            import psycopg2
            from psycopg2.extras import DictCursor
        except ImportError as exc:
            raise RuntimeError("DATABASE_URL is set but psycopg2-binary is missing. Add it to requirements.txt.") from exc
        if dsn.startswith("postgres://"):
            dsn = "postgresql://" + dsn[len("postgres://"):]
        self._driver = psycopg2
        self._conn = psycopg2.connect(dsn, cursor_factory=DictCursor, sslmode="require")
    def execute(self, sql, params=()):
        cur = self._conn.cursor()
        if sql.strip().upper().startswith("PRAGMA TABLE_INFO("):
            table = re.search(r"PRAGMA\s+table_info\(([^)]+)\)", sql, re.I).group(1).strip().strip('"`')
            cur.execute("SELECT ordinal_position-1 AS cid, column_name AS name, data_type AS type, 0 AS not_null, column_default AS dflt_value, 0 AS pk FROM information_schema.columns WHERE table_schema=current_schema() AND table_name=%s ORDER BY ordinal_position", (table,))
            return cur
        cur.execute(sql.replace("?", "%s"), tuple(params or ()))
        return cur
    def executescript(self, script):
        # DDL statements in our schema contain no embedded semicolons.
        for statement in script.split(";"):
            statement = statement.strip()
            if statement:
                statement = re.sub(r"INTEGER\s+PRIMARY\s+KEY\s+AUTOINCREMENT", "SERIAL PRIMARY KEY", statement, flags=re.I)
                self.execute(statement)
    def commit(self): self._conn.commit()
    def rollback(self): self._conn.rollback()
    def close(self): self._conn.close()


def get_db():
    if DATABASE_URL:
        return PostgresConnection(DATABASE_URL)
    con = sqlite3.connect(DATABASE, timeout=20)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys = ON")
    return con

def initialize_database():
    con = get_db()
    con.executescript("""
    CREATE TABLE IF NOT EXISTS users(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT NOT NULL,
        email TEXT UNIQUE NOT NULL,
        password TEXT NOT NULL,
        role TEXT NOT NULL DEFAULT 'student',
        room TEXT DEFAULT '',
        mobile TEXT DEFAULT '',
        student_id TEXT DEFAULT '',
        department TEXT DEFAULT '',
        year TEXT DEFAULT '',
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    CREATE TABLE IF NOT EXISTS workers(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT NOT NULL,
        department TEXT NOT NULL,
        mobile TEXT NOT NULL,
        active INTEGER DEFAULT 1,
        email TEXT UNIQUE,
        password TEXT DEFAULT '',
        worker_code TEXT DEFAULT ''
    );
    CREATE TABLE IF NOT EXISTS complaints(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        complaint_id TEXT UNIQUE NOT NULL,
        user_id INTEGER NOT NULL,
        category TEXT NOT NULL,
        title TEXT NOT NULL,
        description TEXT NOT NULL,
        location TEXT NOT NULL,
        priority TEXT NOT NULL,
        image TEXT DEFAULT '',
        status TEXT NOT NULL DEFAULT 'Submitted',
        assigned_to TEXT DEFAULT '',
        assigned_worker_id INTEGER DEFAULT NULL,
        admin_remark TEXT DEFAULT '',
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY(user_id) REFERENCES users(id),
        FOREIGN KEY(assigned_worker_id) REFERENCES workers(id)
    );
    CREATE TABLE IF NOT EXISTS feedback(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        complaint_id INTEGER NOT NULL,
        rating INTEGER NOT NULL,
        comment TEXT DEFAULT '',
        FOREIGN KEY(complaint_id) REFERENCES complaints(id)
    );
    CREATE TABLE IF NOT EXISTS password_reset_requests(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        account_type TEXT NOT NULL,
        account_id INTEGER NOT NULL,
        email TEXT NOT NULL,
        pending_password_hash TEXT NOT NULL,
        status TEXT NOT NULL DEFAULT 'Pending',
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        reviewed_at TIMESTAMP
    );
    CREATE TABLE IF NOT EXISTS notices(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        title TEXT NOT NULL,
        body TEXT NOT NULL,
        category TEXT NOT NULL DEFAULT 'General',
        priority TEXT NOT NULL DEFAULT 'Normal',
        created_by INTEGER,
        active INTEGER NOT NULL DEFAULT 1,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    CREATE TABLE IF NOT EXISTS notifications(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        account_type TEXT NOT NULL,
        account_id INTEGER NOT NULL,
        title TEXT NOT NULL,
        message TEXT NOT NULL,
        link TEXT DEFAULT '',
        is_read INTEGER NOT NULL DEFAULT 0,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    CREATE TABLE IF NOT EXISTS lost_found(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        user_id INTEGER NOT NULL,
        item_type TEXT NOT NULL,
        title TEXT NOT NULL,
        description TEXT NOT NULL,
        location TEXT NOT NULL,
        item_date TEXT DEFAULT '',
        category TEXT DEFAULT 'Other',
        image_path TEXT DEFAULT '',
        status TEXT NOT NULL DEFAULT 'Open',
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    CREATE TABLE IF NOT EXISTS campus_locations(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT NOT NULL,
        category TEXT NOT NULL,
        building TEXT DEFAULT '',
        description TEXT DEFAULT '',
        hours TEXT DEFAULT '',
        phone TEXT DEFAULT '',
        map_url TEXT DEFAULT '',
        active INTEGER NOT NULL DEFAULT 1,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    CREATE TABLE IF NOT EXISTS ai_chat_messages(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        account_type TEXT NOT NULL,
        account_id INTEGER NOT NULL,
        role TEXT NOT NULL,
        message TEXT NOT NULL,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    CREATE INDEX IF NOT EXISTS idx_notifications_account ON notifications(account_type, account_id, is_read);
    CREATE INDEX IF NOT EXISTS idx_lost_found_status ON lost_found(status, created_at);
    CREATE INDEX IF NOT EXISTS idx_ai_chat_account ON ai_chat_messages(account_type, account_id, id);
    CREATE INDEX IF NOT EXISTS idx_notices_active ON notices(active, created_at);
    """)
    # Upgrade older databases safely.
    existing = {r[1] for r in con.execute("PRAGMA table_info(users)").fetchall()}
    for col, typ in [("mobile","TEXT DEFAULT ''"),("student_id","TEXT DEFAULT ''"),("department","TEXT DEFAULT ''"),("year","TEXT DEFAULT ''"),("created_at","TIMESTAMP")]:
        if col not in existing:
            con.execute(f"ALTER TABLE users ADD COLUMN {col} {typ}")
    existing_w = {r[1] for r in con.execute("PRAGMA table_info(workers)").fetchall()}
    for col, typ in [("email", "TEXT"), ("password", "TEXT DEFAULT ''"), ("worker_code", "TEXT DEFAULT ''")]:
        if col not in existing_w:
            con.execute(f"ALTER TABLE workers ADD COLUMN {col} {typ}")
    existing_c = {r[1] for r in con.execute("PRAGMA table_info(complaints)").fetchall()}
    if "assigned_worker_id" not in existing_c:
        con.execute("ALTER TABLE complaints ADD COLUMN assigned_worker_id INTEGER DEFAULT NULL")
    existing_lf = {r[1] for r in con.execute("PRAGMA table_info(lost_found)").fetchall()}
    if "image_path" not in existing_lf:
        con.execute("ALTER TABLE lost_found ADD COLUMN image_path TEXT DEFAULT ''")
    if "category" not in existing_lf:
        con.execute("ALTER TABLE lost_found ADD COLUMN category TEXT DEFAULT 'Other'")
    if "image_data" not in existing_lf:
        con.execute("ALTER TABLE lost_found ADD COLUMN image_data TEXT DEFAULT ''")
    if "image_mime" not in existing_lf:
        con.execute("ALTER TABLE lost_found ADD COLUMN image_mime TEXT DEFAULT ''")
    # Admin bootstrap is controlled through Vercel environment variables.
    # To reset an EXISTING admin once, set ADMIN_RESET_EXISTING_PASSWORD=true,
    # deploy, then immediately change that flag back to false and redeploy.
    admin_email = os.environ.get("ADMIN_EMAIL", "").strip().lower()
    admin_password = os.environ.get("ADMIN_PASSWORD", "")
    admin_name = os.environ.get("ADMIN_NAME", "Campus Administrator").strip() or "Campus Administrator"
    admin_mobile = os.environ.get("ADMIN_MOBILE", "")
    reset_existing_admin = os.environ.get("ADMIN_RESET_EXISTING_PASSWORD", "").strip().lower() in {"1", "true", "yes", "on"}
    if valid_campus_email(admin_email) and admin_password:
        existing_admin = con.execute("SELECT id, role FROM users WHERE email=?", (admin_email,)).fetchone()
        configured_hash = generate_password_hash(admin_password)
        if not existing_admin:
            con.execute(
                "INSERT INTO users(name,email,password,role,mobile,department) VALUES(?,?,?,?,?,?)",
                (admin_name, admin_email, configured_hash, "admin", admin_mobile, "Campus Administration"),
            )
        elif existing_admin["role"] == "admin" and reset_existing_admin:
            con.execute("UPDATE users SET password=? WHERE id=?", (configured_hash, existing_admin["id"]))
    for name, dept, mobile in DEMO_WORKERS:
        if not con.execute("SELECT 1 FROM workers WHERE name=? AND department=?", (name,dept)).fetchone():
            con.execute("INSERT INTO workers(name,department,mobile) VALUES(?,?,?)", (name,dept,mobile))
    # Give old users a student ID if missing.
    rows = con.execute("SELECT id FROM users WHERE role='student' AND (student_id IS NULL OR student_id='')").fetchall()
    for r in rows:
        con.execute("UPDATE users SET student_id=? WHERE id=?", (f"CC{2026000+r['id']}", r['id']))
    con.commit(); con.close()

def allowed_file(filename):
    return "." in filename and filename.rsplit(".",1)[1].lower() in ALLOWED_EXTENSIONS

def current_user():
    if not session.get("user_id"): return None
    con=get_db(); u=con.execute("SELECT * FROM users WHERE id=?",(session["user_id"],)).fetchone(); con.close(); return u

@app.context_processor
def globals_ctx():
    return {"categories":CATEGORIES,"priorities":PRIORITIES,"statuses":STATUSES,"current_user":current_user()}

# ---------------- UI ----------------
SHELL = r'''
<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{{ title or "Campus Care 360" }}</title><link rel="icon" type="image/x-icon" href="{{ url_for('static', filename='favicon.ico') }}"><link rel="apple-touch-icon" href="{{ url_for('static', filename='campus-care-360-logo.png') }}">
<style>
:root{--ink:#12233f;--muted:#71809a;--paper:#f5f7fb;--white:#fff;--blue:#3568ff;--blue2:#6b7cff;--green:#18a873;--orange:#f29a3f;--red:#e95d6a;--line:#e6eaf2;--shadow:0 18px 45px rgba(24,42,75,.08);--radius:22px}*{box-sizing:border-box}body{margin:0;font-family:Inter,ui-sans-serif,system-ui,-apple-system,Segoe UI,Roboto,Arial;color:var(--ink);background:var(--paper)}a{text-decoration:none;color:inherit}button,input,select,textarea{font:inherit}button{cursor:pointer}.app{min-height:100vh;display:flex}.sidebar{position:fixed;left:0;top:0;bottom:0;width:260px;background:#101d35;color:#fff;padding:26px 18px;display:flex;flex-direction:column;z-index:20;overflow:auto}.brand{display:flex;align-items:center;gap:12px;padding:5px 10px 25px}.logo{width:43px;height:43px;border-radius:14px;background:linear-gradient(135deg,#5d7cff,#31c9ff);display:grid;place-items:center;font-weight:900;box-shadow:0 10px 25px rgba(59,120,255,.28)}.brand h2{font-size:19px;margin:0}.brand small{display:block;color:#8491a9;font-size:11px;margin-top:2px}.side-label{font-size:10px;letter-spacing:1.5px;color:#667590;font-weight:800;padding:0 12px;margin:8px 0 10px}.nav{display:grid;gap:6px}.nav a{display:flex;align-items:center;gap:12px;padding:12px 13px;border-radius:14px;color:#9da9bd;font-size:13px;font-weight:700}.nav a:hover,.nav a.active{background:rgba(255,255,255,.08);color:#fff}.nav-icon{width:20px;text-align:center;font-size:16px}.side-bottom{margin-top:auto;padding-top:20px}.user-mini{border:1px solid rgba(255,255,255,.08);background:rgba(255,255,255,.045);border-radius:18px;padding:13px;display:flex;gap:10px;align-items:center}.avatar{width:38px;height:38px;border-radius:12px;background:linear-gradient(135deg,#dbe4ff,#a9c8ff);color:#173263;display:grid;place-items:center;font-weight:850}.user-mini strong{display:block;font-size:13px}.user-mini span{font-size:10px;color:#8290a8}.logout{display:block;text-align:center;color:#9ba7bb;font-size:12px;margin-top:13px}.main{margin-left:260px;min-height:100vh;width:calc(100% - 260px)}.topbar{height:82px;background:rgba(255,255,255,.9);backdrop-filter:blur(14px);border-bottom:1px solid var(--line);display:flex;align-items:center;justify-content:space-between;padding:0 42px;position:sticky;top:0;z-index:10}.eyebrow{font-size:10px;letter-spacing:2px;font-weight:850;color:#8290a7}.top-right{display:flex;gap:10px;align-items:center}.content{padding:38px 42px 60px;max-width:1450px;margin:auto}.flash{margin:0 42px 18px;padding:13px 16px;border-radius:13px;font-size:13px;font-weight:650}.flash.success{background:#e7f8f0;color:#08764c}.flash.error{background:#fff0f1;color:#b52e3e}.page-head{display:flex;justify-content:space-between;gap:25px;align-items:flex-end;margin-bottom:28px}.page-head h1{font-size:34px;line-height:1.05;margin:7px 0 8px;letter-spacing:-1.3px}.page-head p{margin:0;color:var(--muted);font-size:14px}.btn{border:0;border-radius:13px;padding:12px 17px;font-weight:800;display:inline-flex;align-items:center;gap:8px;transition:.2s}.btn:hover{transform:translateY(-1px)}.btn-primary{background:#152742;color:#fff;box-shadow:0 10px 25px rgba(20,40,70,.18)}.btn-blue{background:var(--blue);color:#fff}.btn-light{background:#fff;border:1px solid var(--line);color:var(--ink)}.hero{position:relative;overflow:hidden;min-height:265px;border-radius:28px;padding:34px;color:#fff;background:radial-gradient(circle at 85% 20%,rgba(92,125,255,.8),transparent 26%),linear-gradient(135deg,#13233e,#203f73 58%,#3568b6);box-shadow:var(--shadow);margin-bottom:31px}.hero:before{content:"";position:absolute;width:330px;height:330px;border:1px solid rgba(255,255,255,.13);border-radius:50%;right:-80px;top:-90px}.hero:after{content:"";position:absolute;width:210px;height:210px;border:1px dashed rgba(255,255,255,.18);border-radius:50%;right:15px;top:-20px}.hero-content{position:relative;z-index:2;max-width:610px}.hero-kicker{font-size:10px;letter-spacing:2px;font-weight:850;color:#a9d8ff}.hero h2{font-size:29px;letter-spacing:-1px;margin:9px 0}.hero p{color:#c8d5e8;max-width:560px;line-height:1.6;font-size:13px}.hero-link{display:inline-block;margin-top:13px;background:#fff;color:#162842;padding:11px 16px;border-radius:12px;font-weight:800;font-size:13px}.orbit{position:absolute;right:80px;top:75px;width:88px;height:88px;border-radius:50%;border:2px solid rgba(255,255,255,.35);z-index:2}.orbit:after{content:"✓";position:absolute;right:-8px;top:12px;width:35px;height:35px;border-radius:11px;background:#fff;color:#1e62da;display:grid;place-items:center;font-weight:900;box-shadow:0 10px 30px rgba(0,0,0,.18)}.section-title{display:flex;align-items:center;justify-content:space-between;margin:0 0 15px}.section-title h3{margin:0;font-size:17px}.section-title span{font-size:11px;color:var(--muted)}.stats{display:grid;grid-template-columns:repeat(4,1fr);gap:14px;margin-bottom:31px}.stat{background:#fff;border:1px solid var(--line);border-radius:18px;padding:19px;box-shadow:0 8px 28px rgba(24,42,75,.04)}.stat-top{display:flex;justify-content:space-between;align-items:center;color:var(--muted);font-size:12px;font-weight:700}.stat-icon{width:34px;height:34px;border-radius:11px;background:#eef3ff;display:grid;place-items:center}.stat strong{display:block;font-size:30px;margin-top:9px}.card{background:#fff;border:1px solid var(--line);border-radius:22px;box-shadow:var(--shadow);padding:25px}.form-grid{display:grid;grid-template-columns:1fr 1fr;gap:18px}.full{grid-column:1/-1}label{display:block;font-size:11px;font-weight:850;color:#52617a;margin:0 0 8px}input,select,textarea{width:100%;border:1px solid #e1e6ef;background:#fbfcfe;color:var(--ink);border-radius:13px;padding:12px 13px;outline:none}input:focus,select:focus,textarea:focus{border-color:#7694ff;box-shadow:0 0 0 4px #edf1ff}textarea{min-height:125px;resize:vertical}.form-actions{display:flex;justify-content:flex-end;gap:10px;margin-top:22px}.upload{border:1.5px dashed #cdd6e4;border-radius:15px;padding:18px;text-align:center;background:#fafcff;color:#728098;font-size:12px}.upload input{border:0;background:transparent;padding:0;margin-top:8px}.auth-wrap{min-height:100vh;display:grid;place-items:center;padding:30px;background:radial-gradient(circle at 10% 10%,#e5edff,transparent 32%),#f5f7fb}.auth-card{width:min(480px,100%);background:#fff;border:1px solid var(--line);border-radius:28px;padding:34px;box-shadow:var(--shadow)}.auth-brand{display:flex;align-items:center;gap:11px;margin-bottom:25px}.auth-brand h1{font-size:20px;margin:0}.auth-brand span{display:block;color:var(--muted);font-size:11px;margin-top:2px}.auth-card h2{font-size:29px;margin:0 0 7px}.auth-card>p{color:var(--muted);font-size:13px;margin:0 0 23px}.auth-card form{display:grid;gap:14px}.auth-link{text-align:center;font-size:12px;color:var(--muted);margin-top:18px}.auth-link a{color:var(--blue);font-weight:800}.demo{margin-top:22px;background:#f3f6fb;border-radius:15px;padding:13px;font-size:11px;color:#63718a;line-height:1.6}.demo b{color:var(--ink)}.complaints{display:grid;gap:10px}.complaint-row{background:#fff;border:1px solid var(--line);border-radius:18px;padding:16px 18px;display:grid;grid-template-columns:1fr auto;gap:14px;align-items:center}.row-meta{font-size:10px;color:#8995a9;font-weight:750;margin-bottom:6px}.row-title{font-size:14px;font-weight:820}.row-sub{font-size:11px;color:var(--muted);margin-top:6px}.row-right{display:flex;align-items:center;gap:12px}.chip{display:inline-flex;align-items:center;gap:5px;padding:6px 9px;border-radius:99px;font-size:10px;font-weight:850;white-space:nowrap}.chip.submitted{background:#eef3ff;color:#4168cc}.chip.assigned{background:#f4edff;color:#8054bc}.chip.progress{background:#fff4df;color:#ad6b12}.chip.resolved{background:#e8f8f0;color:#08794e}.chip.closed{background:#edf0f3;color:#657184}.chip.emergency{background:#ffe7ea;color:#c53a49}.arrow{color:#99a5b8;font-size:19px}.detail-grid{display:grid;grid-template-columns:1.55fr .8fr;gap:20px}.detail-title{display:flex;justify-content:space-between;gap:15px;align-items:flex-start}.detail-title h1{margin:0;font-size:27px}.mini-meta{font-size:11px;color:var(--muted);margin-top:8px}.detail-block{margin-top:23px}.detail-block h4{font-size:11px;text-transform:uppercase;letter-spacing:1px;color:#8390a5;margin:0 0 9px}.detail-block p{font-size:13px;line-height:1.65;color:#516078;margin:0}.photo{width:100%;max-height:300px;object-fit:cover;border-radius:16px}.timeline{display:grid}.step{display:grid;grid-template-columns:28px 1fr;gap:11px;position:relative;min-height:58px}.step:not(:last-child):before{content:"";position:absolute;left:10px;top:21px;bottom:0;width:2px;background:#e5e9f1}.dot{width:22px;height:22px;border-radius:50%;background:#e8edf5;color:#8a96aa;display:grid;place-items:center;font-size:10px;font-weight:900;z-index:2}.step.done .dot{background:#e4f7ee;color:#0b8b5a}.step.current .dot{background:#e8eeff;color:#3867dd}.step b{font-size:12px}.step span{display:block;color:#8b97a9;font-size:10px;margin-top:3px}.info-list{display:grid;gap:13px}.info-item{display:flex;justify-content:space-between;gap:15px;font-size:12px}.info-item span{color:var(--muted)}.info-item strong{text-align:right}.feedback{margin-top:18px;background:#f8fafc;border:1px solid var(--line);border-radius:17px;padding:17px}.toolbar{background:#fff;border:1px solid var(--line);border-radius:18px;padding:12px;display:grid;grid-template-columns:1fr 190px auto;gap:10px;margin-bottom:18px}.admin-list{display:grid;gap:14px}.admin-card{background:#fff;border:1px solid var(--line);border-radius:20px;padding:19px;box-shadow:0 8px 28px rgba(24,42,75,.04)}.admin-top{display:flex;justify-content:space-between;gap:15px}.admin-title{font-weight:850;font-size:15px;margin:5px 0}.admin-meta{font-size:10px;color:#8490a4}.admin-desc{font-size:12px;color:#637089;line-height:1.55;margin:13px 0}.admin-tags{display:flex;gap:7px;flex-wrap:wrap}.admin-form{margin-top:15px;padding-top:15px;border-top:1px solid #edf0f5;display:grid;grid-template-columns:1fr 1fr 1.4fr auto;gap:8px;align-items:end}.table-wrap{overflow:auto;background:#fff;border:1px solid var(--line);border-radius:20px;box-shadow:var(--shadow)}table{width:100%;border-collapse:collapse;min-width:850px}th,td{text-align:left;padding:14px 15px;border-bottom:1px solid #edf0f5;font-size:12px}th{font-size:10px;text-transform:uppercase;letter-spacing:.8px;color:#8490a4;background:#fafbfd}td strong{font-size:12px}.searchbar{display:grid;grid-template-columns:1fr auto;gap:10px;margin-bottom:17px}.profile-grid{display:grid;grid-template-columns:1fr 1fr;gap:20px}.contact-grid,.worker-grid,.help-grid{display:grid;grid-template-columns:repeat(3,1fr);gap:15px}.contact-card,.worker-card,.help-card{background:#fff;border:1px solid var(--line);border-radius:20px;padding:20px;box-shadow:0 8px 28px rgba(24,42,75,.04)}.contact-card h3,.worker-card h3,.help-card h3{margin:0 0 6px;font-size:15px}.muted{color:var(--muted);font-size:12px;line-height:1.6}.phone{display:inline-block;margin-top:12px;font-weight:850;color:#315fda}.faq{display:grid;gap:10px}.faq details{background:#fff;border:1px solid var(--line);border-radius:16px;padding:15px}.faq summary{font-weight:800;font-size:13px;cursor:pointer}.faq p{color:var(--muted);font-size:12px;line-height:1.6}.ai-box{display:grid;grid-template-rows:1fr auto;height:520px;background:#fff;border:1px solid var(--line);border-radius:24px;overflow:hidden;box-shadow:var(--shadow)}.chat{padding:20px;overflow:auto;background:linear-gradient(#fbfcff,#f5f7fb)}.msg{max-width:78%;padding:12px 14px;border-radius:16px;margin-bottom:10px;font-size:12px;line-height:1.55}.msg.ai{background:#fff;border:1px solid var(--line)}.msg.user{background:#152742;color:#fff;margin-left:auto}.chat-form{display:flex;gap:8px;padding:13px;border-top:1px solid var(--line);background:#fff}.chat-form input{flex:1}.floating-ai{position:fixed;right:25px;bottom:24px;width:58px;height:58px;border-radius:50%;background:linear-gradient(135deg,#3568ff,#6b7cff);color:#fff;display:grid;place-items:center;font-size:23px;box-shadow:0 15px 35px rgba(53,104,255,.35);z-index:15}.empty{background:#fff;border:1px dashed #ccd5e3;border-radius:18px;padding:35px;text-align:center;color:#7c899d;font-size:13px}.mobile-menu{display:none}
@media(max-width:1050px){.sidebar{width:220px}.main{margin-left:220px;width:calc(100% - 220px)}.stats{grid-template-columns:repeat(2,1fr)}.detail-grid,.profile-grid{grid-template-columns:1fr}.contact-grid,.worker-grid,.help-grid{grid-template-columns:repeat(2,1fr)}.admin-form{grid-template-columns:1fr 1fr}.toolbar{grid-template-columns:1fr 180px}.orbit{right:35px}}
@media(max-width:760px){.sidebar{position:fixed;transform:translateX(-100%);transition:.25s;width:250px}.sidebar.open{transform:translateX(0)}.main{margin-left:0;width:100%}.topbar{padding:0 18px}.content{padding:25px 18px 45px}.flash{margin:0 18px 15px}.mobile-menu{display:block;border:0;background:#f3e9d9;border-radius:10px;width:38px;height:38px}.page-head{align-items:flex-start;flex-direction:column}.page-head h1{font-size:28px}.hero{padding:25px}.orbit{opacity:.35}.stats{grid-template-columns:1fr 1fr}.complaint-row{grid-template-columns:1fr}.row-right{justify-content:space-between}.toolbar,.searchbar{grid-template-columns:1fr}.admin-form,.form-grid{grid-template-columns:1fr}.full{grid-column:auto}.contact-grid,.worker-grid,.help-grid{grid-template-columns:1fr}.ai-box{height:600px}.auth-card{padding:25px}}
/* CAMPUS CARE SIGNATURE THEME: oxblood red, cream, porcelain white and near-black */
:root{--ink:#171311;--muted:#746b63;--paper:#f4eee5;--white:#fffdf9;--blue:#a51f2b;--blue2:#d44a4f;--green:#24734e;--orange:#bd7b32;--red:#b51f2c;--line:#e5d9cc;--shadow:0 22px 60px rgba(38,20,15,.12);--radius:22px}
body{background:radial-gradient(circle at 88% 0%,rgba(165,31,43,.07),transparent 27%),var(--paper);color:var(--ink)}
.sidebar{background:linear-gradient(165deg,#171311,#2a1717 70%,#511a20);border-right:1px solid rgba(255,255,255,.08)}
.logo{background:linear-gradient(135deg,#a51f2b,#e8b8a5);color:#fff;box-shadow:0 10px 28px rgba(165,31,43,.32);border-radius:13px}
.nav a:hover,.nav a.active{background:rgba(214,72,76,.22);color:#fff;box-shadow:inset 3px 0 #d84b50}
.topbar{background:rgba(255,253,249,.91);border-color:var(--line)}
.eyebrow{color:#9b2630}.btn-primary{background:#171311;color:#fff;box-shadow:0 10px 25px rgba(23,19,17,.2)}.btn-blue{background:linear-gradient(135deg,#a51f2b,#c83b43);color:#fff;box-shadow:0 10px 24px rgba(165,31,43,.18)}.btn-light{background:#fffdf9;border-color:#dfd0c2;color:#241a17}
.hero{background:radial-gradient(circle at 85% 15%,rgba(236,190,161,.28),transparent 27%),linear-gradient(125deg,#171311 0%,#35191b 54%,#8c2029 100%);border:1px solid rgba(255,255,255,.14);box-shadow:0 24px 60px rgba(67,20,22,.2)}.hero-kicker{color:#efc6ad}.hero p{color:#eadbd0}.hero-link{color:#241614;background:#fff5e9}.orbit{border-color:rgba(245,225,207,.5)}.orbit:after{color:#a51f2b}
.stat,.card,.admin-card,.contact-card,.worker-card,.help-card,.complaint-row,.table-wrap{background:#fffdf9;border-color:var(--line);box-shadow:0 12px 34px rgba(39,24,17,.055)}
.stat-icon,.avatar{background:linear-gradient(135deg,#f1ded0,#e7b7a5);color:#761923}.chip.submitted{background:#f7e7e4;color:#92212b}.chip.assigned{background:#f3e6d6;color:#83501d}.chip.progress{background:#f7eed7;color:#8e5c13}.chip.resolved{background:#e5f1e8;color:#246443}.chip.closed{background:#ece7e1;color:#615850}.chip.emergency{background:#f9dfe0;color:#9f1422}
input,select,textarea{background:#fffdf9;border-color:#dfd2c6}input:focus,select:focus,textarea:focus{border-color:#b83a42;box-shadow:0 0 0 4px rgba(165,31,43,.1)}.auth-wrap{background:radial-gradient(circle at 15% 15%,rgba(165,31,43,.13),transparent 28%),radial-gradient(circle at 90% 85%,rgba(206,157,124,.22),transparent 25%),#f4eee5}.auth-card{background:rgba(255,253,249,.96);border-color:#e4d5c7;box-shadow:0 28px 80px rgba(52,23,18,.14)}.auth-link a,.phone{color:#a51f2b}.demo{background:#f5eae0;color:#6b5146}.floating-ai{background:linear-gradient(135deg,#a51f2b,#171311)}
.hero:before{border-color:rgba(255,239,224,.18)}.hero:after{border-color:rgba(255,239,224,.24)}.flash.success{background:#e5f1e8;color:#246443}.flash.error{background:#f9e2e1;color:#921b25}


.login-emergency{display:grid;grid-template-columns:1fr;gap:9px;margin-top:18px}.login-sos{display:flex;align-items:center;justify-content:center;padding:12px;border-radius:12px;background:#a51f30;color:#fff;font-weight:950;font-size:12px;letter-spacing:.3px}.login-emergency-link{display:block;text-align:center;border:1px solid #e7dcd2;border-radius:12px;padding:11px;color:#8d1b2a;font-size:12px;font-weight:850}.login-sos:hover{background:#741421}

/* Campus Care premium identity: oxblood, cream, ivory and charcoal */
:root{--ink:#171313;--muted:#756b66;--paper:#f6f0e8;--white:#fffdf9;--blue:#9f1d2d;--blue2:#c3424d;--green:#27785a;--orange:#b86a28;--red:#b51f32;--line:#e7dcd2;--shadow:0 20px 55px rgba(42,18,18,.10);--radius:22px}
body{background:var(--paper);color:var(--ink)}
.sidebar{background:#171313;border-right:1px solid #342326}
.logo{width:46px;height:46px;border-radius:15px;background:linear-gradient(145deg,#b51f32,#711421);color:#fff;display:flex;align-items:center;justify-content:center;gap:1px;font-size:13px;letter-spacing:-.8px;font-weight:950;box-shadow:0 10px 28px rgba(159,29,45,.32);border:1px solid rgba(255,255,255,.22)}
.logo span{font-size:21px;letter-spacing:-2px}
.btn-primary,.btn-blue{background:#941b2b;color:#fff;box-shadow:0 10px 25px rgba(148,27,43,.20)}
.btn-primary:hover,.btn-blue:hover{background:#741421}
.hero{background:radial-gradient(circle at 86% 15%,rgba(190,73,83,.45),transparent 30%),linear-gradient(125deg,#171313,#3b171e 62%,#791a29)}
.hero-kicker{color:#f1c8be}.hero p{color:#f0ded8}.hero-link{color:#6f1522;background:#fffaf3}
.stat,.card,.admin-card,.contact-card,.worker-card,.help-card,.complaint-row,.table-wrap{background:#fffdf9;border-color:#e7dcd2}
.stat-icon{background:#f8e8e5;color:#9f1d2d}
.nav a:hover,.nav a.active{background:rgba(181,31,50,.24);color:#fff}
.auth-wrap{background:radial-gradient(circle at 8% 10%,#ead4c9,transparent 34%),#f6f0e8}
.auth-card{background:#fffdf9;border-color:#e7dcd2}
.auth-link a,.phone{color:#9f1d2d}
input:focus,select:focus,textarea:focus{border-color:#b44b55;box-shadow:0 0 0 4px #f8e6e3}
.emergency-actions{display:grid;grid-template-columns:1.1fr 1fr;gap:14px;margin:0 0 28px}
.emergency-tile{display:flex;align-items:center;gap:14px;padding:19px 21px;border-radius:20px;border:1px solid #ead5d0;background:#fffdf9;box-shadow:0 10px 28px rgba(42,18,18,.05)}
.emergency-tile.sos{background:linear-gradient(125deg,#a51f30,#741421);color:#fff;border-color:#a51f30}
.emergency-tile .tile-icon{width:44px;height:44px;display:grid;place-items:center;border-radius:14px;background:#f8e6e3;color:#a51f30;font-size:22px;font-weight:950;flex-shrink:0}
.emergency-tile.sos .tile-icon{background:rgba(255,255,255,.15);color:#fff}
.emergency-tile strong{display:block;font-size:14px}.emergency-tile small{display:block;margin-top:4px;font-size:11px;opacity:.78;line-height:1.4}
.emergency-tile .tile-arrow{margin-left:auto;font-size:22px;font-weight:800}
.sos-menu{color:#ffd9d5!important}.sos-menu .nav-icon{color:#ff9ea5}
.emergency-panel{border-radius:24px;background:linear-gradient(135deg,#1a1415,#481923);color:#fff;padding:28px;margin-bottom:20px}
.emergency-panel h2{margin:8px 0;font-size:27px}.emergency-panel p{color:#e7cfd0;line-height:1.6;font-size:13px}
.emergency-number{font-size:34px;font-weight:950;letter-spacing:-1px;margin:10px 0}
.emergency-call{display:inline-flex;align-items:center;gap:9px;background:#b51f32;color:white;padding:13px 18px;border-radius:13px;font-weight:900}
@media(max-width:760px){.emergency-actions{grid-template-columns:1fr}.emergency-tile{padding:15px}.emergency-panel{padding:22px}}


/* ChatGPT-inspired Campus Care 360 AI workspace */
.chatgpt-shell{min-height:calc(100vh - 160px);height:calc(100vh - 160px);max-height:980px;display:flex;flex-direction:column;background:#fffdf9;border:1px solid #e9e2d8;border-radius:24px;overflow:hidden;box-shadow:0 18px 55px rgba(30,20,18,.07);margin-top:-12px}
.chatgpt-header{display:flex;align-items:center;justify-content:space-between;padding:18px 24px;border-bottom:1px solid #eee7df;background:rgba(255,253,249,.96)}
.chatgpt-title-wrap{display:flex;align-items:center;gap:12px}.chatgpt-mark,.chatgpt-avatar{width:40px;height:40px;display:grid;place-items:center;border-radius:14px;background:linear-gradient(135deg,#a61e2d,#401016);color:#fff;font-size:20px;box-shadow:0 6px 18px rgba(126,25,39,.2)}
.chatgpt-title-wrap h1{font-size:16px;margin:0;color:#211b19;letter-spacing:-.3px}.chatgpt-title-wrap p{margin:3px 0 0;color:#91857d;font-size:11px}.chatgpt-clear-form{margin:0}.chatgpt-clear{border:1px solid #e6ddd4;border-radius:12px;background:#fff;color:#3b302b;padding:10px 13px;font-size:12px;font-weight:750;display:flex;gap:7px;align-items:center}.chatgpt-clear:hover{background:#f8f1e9}
.chatgpt-conversation{flex:1;overflow:auto;padding:28px clamp(18px,7vw,100px);scroll-behavior:smooth;background:radial-gradient(ellipse at 50% 0%,rgba(248,238,225,.55),transparent 46%),#fffdf9}.chatgpt-welcome{max-width:780px;margin:5vh auto 20px;text-align:center}.welcome-orb{width:64px;height:64px;margin:0 auto 22px;display:grid;place-items:center;border-radius:22px;color:white;font-size:30px;background:linear-gradient(145deg,#a91e2c,#261c1c);box-shadow:0 12px 35px rgba(122,25,34,.2)}.chatgpt-welcome h2{font-size:clamp(25px,3vw,35px);letter-spacing:-1.2px;color:#201a18;margin:0 0 12px}.chatgpt-welcome>p{max-width:510px;margin:0 auto;color:#887c74;font-size:14px;line-height:1.7}.chatgpt-suggestions{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:12px;margin:32px auto 0;text-align:left;max-width:670px}.chatgpt-suggestions button{position:relative;display:grid;grid-template-columns:30px 1fr;grid-template-rows:auto auto;column-gap:10px;align-items:center;text-align:left;padding:16px;border:1px solid #eee5dc;background:rgba(255,255,255,.85);border-radius:17px;transition:transform .18s,border-color .18s,box-shadow .18s}.chatgpt-suggestions button:hover{transform:translateY(-2px);border-color:#d5b9ae;box-shadow:0 9px 24px rgba(44,26,20,.06)}.chatgpt-suggestions button span{grid-row:1/3;width:30px;height:30px;display:grid;place-items:center;border-radius:10px;background:#f8eee8;color:#922431;font-size:17px}.chatgpt-suggestions button b{font-size:12px;color:#352a26}.chatgpt-suggestions button small{font-size:11px;color:#92857c;margin-top:4px}
.chatgpt-message{max-width:850px;margin:0 auto 25px;display:flex;gap:13px;align-items:flex-start;line-height:1.75;font-size:14px;white-space:pre-wrap;overflow-wrap:anywhere}.chatgpt-message.from-user{justify-content:flex-end}.chatgpt-message.from-user .chatgpt-message-content{max-width:min(80%,650px);background:#f3eee8;color:#2b2420;padding:12px 16px;border-radius:19px 19px 5px 19px}.chatgpt-message.from-ai{max-width:850px}.chatgpt-message.from-ai .chatgpt-avatar{flex:0 0 32px;width:32px;height:32px;border-radius:11px;font-size:15px}.chatgpt-message-content{min-width:0;max-width:calc(100% - 45px);color:#352d29}.chatgpt-speaker{font-size:12px;font-weight:850;color:#8c202b;margin:2px 0 5px}.chatgpt-message-text{font-size:14px;line-height:1.8}.chatgpt-composer-wrap{padding:12px clamp(14px,6vw,78px) 13px;background:linear-gradient(180deg,rgba(255,253,249,.8),#fffdf9 30%)}.chatgpt-composer{max-width:850px;margin:0 auto;display:flex;align-items:flex-end;gap:12px;padding:10px 10px 10px 18px;border:1px solid #e6dcd2;border-radius:24px;background:white;box-shadow:0 8px 28px rgba(35,22,18,.08);transition:border-color .2s,box-shadow .2s}.chatgpt-composer:focus-within{border-color:#c89a91;box-shadow:0 0 0 4px rgba(160,36,48,.07),0 8px 28px rgba(35,22,18,.06)}.chatgpt-composer textarea{border:0!important;box-shadow:none!important;outline:none;resize:none;min-height:28px;max-height:150px;padding:5px 0!important;background:transparent!important;border-radius:0!important;font-size:14px;line-height:1.55;color:#2c2421}.chatgpt-composer textarea::placeholder{color:#a1968e}.chatgpt-send{flex:0 0 38px;width:38px;height:38px;border:0;border-radius:14px;background:#8e1e2a;color:white;font-size:23px;font-weight:800;display:grid;place-items:center;transition:background .2s,transform .2s}.chatgpt-send:hover{background:#63141e;transform:translateY(-1px)}.chatgpt-disclaimer{text-align:center;font-size:10px;color:#a1968e;margin:10px 0 0}.chatgpt-disclaimer a{color:#8e1e2a;font-weight:850;text-decoration:underline}
@media(max-width:760px){.chatgpt-shell{height:calc(100vh - 125px);min-height:570px;margin:0 -5px;border-radius:18px}.chatgpt-header{padding:14px}.chatgpt-clear span{display:none}.chatgpt-conversation{padding:22px 14px}.chatgpt-welcome{margin:5vh auto 15px}.chatgpt-suggestions{grid-template-columns:1fr;margin-top:24px}.chatgpt-suggestions button{padding:13px}.chatgpt-message{font-size:13px;gap:9px}.chatgpt-message.from-user .chatgpt-message-content{max-width:90%}.chatgpt-composer-wrap{padding:10px 10px 12px}.chatgpt-composer{padding-left:14px}}


/* 2026 visual refresh: signature mark, help centre, profiles and chat polish */
.logo.logo-mark{position:relative;overflow:hidden;gap:0;width:48px;height:48px;border-radius:16px;background:linear-gradient(145deg,#c82b3e 0%,#8e1728 55%,#3a1118 100%);box-shadow:0 12px 28px rgba(142,23,40,.28);border:1px solid rgba(255,255,255,.32);isolation:isolate}
.logo.logo-mark:before{content:"";position:absolute;inset:5px;border:1px solid rgba(255,255,255,.24);border-radius:12px;transform:rotate(-8deg);z-index:-1}
.logo.logo-mark:after{content:"";position:absolute;width:23px;height:23px;border:2px solid rgba(255,255,255,.24);border-radius:50%;right:-8px;top:-7px}
.logo.logo-mark{display:flex;align-items:center;justify-content:center;gap:0}
.logo-c{font-size:22px!important;font-weight:950;letter-spacing:-1.5px!important;line-height:1;color:#E8D29A}
.logo-degree{font-size:9px;font-weight:950;letter-spacing:-.2px;color:#E8D29A;align-self:center;margin:0 0 0 -1px;white-space:nowrap}
.chatgpt-shell{border-radius:28px;box-shadow:0 24px 70px rgba(51,24,24,.09);border-color:#e5d7cb}
.chatgpt-header{background:linear-gradient(110deg,#fffdf9,#fbf1e9)}
.chatgpt-mark,.chatgpt-avatar,.welcome-orb{background:linear-gradient(145deg,#cf3547,#741421 72%,#241417);box-shadow:0 10px 25px rgba(142,23,40,.23)}
.chatgpt-send{background:linear-gradient(145deg,#b62436,#781421);box-shadow:0 5px 13px rgba(142,23,40,.18)}
.chatgpt-suggestions button{background:linear-gradient(135deg,#fffdf9,#fff8f1)}
.chatgpt-suggestions button span{background:#f7e5df;color:#9c1d2d}
.chatgpt-message.from-user .chatgpt-message-content{background:#f5e8e0;border:1px solid #eddbd0}
.chatgpt-message.from-ai .chatgpt-message-content{padding:1px 0}
.chatgpt-message-text{color:#332825}
.help-hero{position:relative;overflow:hidden;border-radius:26px;padding:30px;background:radial-gradient(circle at 90% 15%,rgba(208,66,78,.42),transparent 27%),linear-gradient(125deg,#171313,#421820 65%,#8e1c2b);color:#fff;margin-bottom:22px;box-shadow:0 20px 48px rgba(58,20,26,.14)}
.help-hero:after{content:"?";position:absolute;right:6%;top:-38px;font-size:190px;line-height:1;font-weight:950;color:rgba(255,255,255,.055);pointer-events:none}
.help-hero h2{font-size:clamp(23px,3vw,32px);letter-spacing:-1px;margin:8px 0}.help-hero p{color:#f0d9d6;max-width:640px;font-size:13px;line-height:1.75;margin:0}
.help-quicklinks{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:13px;margin:0 0 24px}
.help-quicklink{display:flex;gap:12px;align-items:center;background:#fffdf9;border:1px solid #e7dcd2;border-radius:18px;padding:16px;box-shadow:0 8px 24px rgba(42,18,18,.04);transition:transform .18s,box-shadow .18s}
.help-quicklink:hover{transform:translateY(-3px);box-shadow:0 14px 30px rgba(42,18,18,.09)}
.help-quick-icon{width:42px;height:42px;display:grid;place-items:center;flex-shrink:0;border-radius:14px;background:#f8e7e2;color:#9f1d2d;font-size:20px}
.help-quicklink b{display:block;font-size:12px}.help-quicklink small{display:block;color:#887c74;font-size:10px;line-height:1.5;margin-top:4px}
.help-grid{grid-template-columns:repeat(3,minmax(0,1fr));gap:14px}.help-card{position:relative;overflow:hidden;padding:22px;transition:transform .18s,border-color .18s,box-shadow .18s}.help-card:hover{transform:translateY(-3px);border-color:#d7b8b0;box-shadow:0 16px 34px rgba(42,18,18,.08)}
.help-card .help-card-icon{width:44px;height:44px;border-radius:15px;background:#f8e8e3;color:#9f1d2d;display:grid;place-items:center;font-size:21px;margin-bottom:15px}.help-card h3{font-size:14px}.help-card .muted{margin-bottom:0}
.faq details{box-shadow:0 5px 18px rgba(42,18,18,.025);border-color:#e7dcd2}.faq summary{list-style:none;display:flex;justify-content:space-between;gap:15px}.faq summary:after{content:"＋";color:#9f1d2d;font-weight:950}.faq details[open] summary:after{content:"−"}.faq p{line-height:1.8;margin-bottom:2px}
.profile-hero{position:relative;overflow:hidden;display:flex;align-items:center;gap:20px;padding:27px;border-radius:25px;background:radial-gradient(circle at 90% 15%,rgba(213,82,89,.4),transparent 25%),linear-gradient(125deg,#171313,#431820 62%,#8c1c2b);color:#fff;margin-bottom:20px;box-shadow:0 20px 48px rgba(58,20,26,.12)}
.profile-hero .profile-avatar{width:82px;height:82px;flex:0 0 82px;border-radius:25px;display:grid;place-items:center;background:linear-gradient(145deg,#fff3e7,#e4b6a8);color:#781421;font-size:31px;font-weight:950;border:2px solid rgba(255,255,255,.55);box-shadow:0 12px 28px rgba(0,0,0,.14)}
.profile-hero h1{font-size:clamp(23px,3vw,31px);margin:0 0 6px;letter-spacing:-.7px}.profile-hero p{margin:0;color:#f0d9d6;font-size:12px;line-height:1.7}.profile-role{display:inline-flex;margin-top:11px;padding:6px 10px;border-radius:99px;background:rgba(255,255,255,.13);border:1px solid rgba(255,255,255,.18);font-size:10px;font-weight:850;letter-spacing:.5px}
.profile-info-card{padding:23px}.profile-info-card h3{margin:0 0 18px;font-size:15px}.profile-details{display:grid;grid-template-columns:1fr 1fr;gap:12px}.profile-detail{padding:14px;border:1px solid #eee2d8;background:#fffaf4;border-radius:15px;min-width:0}.profile-detail span{display:block;color:#8b7e75;font-size:10px;font-weight:800;text-transform:uppercase;letter-spacing:.7px;margin-bottom:7px}.profile-detail strong{display:block;font-size:12px;overflow-wrap:anywhere;color:#302522}.profile-stat-card{position:relative;overflow:hidden;background:linear-gradient(145deg,#fffdf9,#f8ece4);border:1px solid #e8d8ce;border-radius:20px;padding:20px}.profile-stat-card strong{font-size:32px;display:block;margin-top:10px;color:#8e1c2b}.profile-stat-card small{color:#83766e;font-size:11px}
.profile-shortcuts{display:grid;gap:10px;margin-top:14px}.profile-shortcut{display:flex;align-items:center;gap:11px;border:1px solid #e8dcd2;background:#fffdf9;border-radius:14px;padding:13px;font-size:12px;font-weight:800}.profile-shortcut span{width:33px;height:33px;display:grid;place-items:center;border-radius:11px;background:#f7e5df;color:#9f1d2d}.profile-shortcut em{font-style:normal;margin-left:auto;color:#a39890;font-size:17px}
@media(max-width:1050px){.help-grid{grid-template-columns:repeat(2,minmax(0,1fr))}.help-quicklinks{grid-template-columns:repeat(3,minmax(0,1fr))}}
@media(max-width:760px){.help-grid,.help-quicklinks{grid-template-columns:1fr}.help-hero{padding:23px}.help-hero:after{right:-5px;font-size:130px}.profile-hero{padding:21px;align-items:flex-start;gap:14px}.profile-hero .profile-avatar{width:64px;height:64px;flex-basis:64px;border-radius:20px;font-size:24px}.profile-details{grid-template-columns:1fr}.profile-info-card{padding:18px}}


/* Accessible show/hide password control */
.password-wrap{position:relative;width:100%;min-width:0}
.password-wrap>input[type=password],.password-wrap>input[type=text]{width:100%;padding-right:76px}
.password-toggle{position:absolute;right:9px;top:50%;transform:translateY(-50%);border:0;background:transparent;color:#941b2b;font-size:12px;font-weight:850;padding:7px 5px;border-radius:7px;cursor:pointer;z-index:2}
.password-toggle:hover{background:#f8e8e5}
.password-toggle:focus-visible{outline:2px solid #941b2b;outline-offset:2px}
/* Use the exact user-provided Campus Care 360 logo asset */
.brand-logo-img{display:block;width:48px;height:48px;flex:0 0 48px;object-fit:contain;border-radius:14px}
.auth-brand .brand-logo-img{width:48px;height:48px;flex-basis:48px}
</style></head><body>
{% if session.get('user_id') or session.get('worker_id') %}<div class="app"><aside class="sidebar" id="sidebar"><div class="brand"><img class="brand-logo-img" src="{{ url_for('static', filename='campus-care-360-logo.png') }}" alt="Campus Care 360 logo"><div><h2>Campus Care 360</h2><small>Campus support hub</small></div></div><div class="side-label">{{ 'ADMIN SPACE' if session.get('role')=='admin' else ('WORKER SPACE' if session.get('role')=='worker' else 'STUDENT SPACE') }}</div><nav class="nav">
<a href="{{url_for('emergency_help')}}" class="sos-menu"><span class="nav-icon">🆘</span>Emergency / SOS</a>
{% if session.get('role')=='admin' %}<a href="{{url_for('admin_dashboard')}}" class="{{'active' if request.endpoint=='admin_dashboard' else ''}}"><span class="nav-icon">▦</span>Dashboard</a><a href="{{url_for('admin_students')}}" class="{{'active' if request.endpoint=='admin_students' else ''}}"><span class="nav-icon">👨‍🎓</span>Students</a><a href="{{url_for('admin_workers')}}" class="{{'active' if request.endpoint=='admin_workers' else ''}}"><span class="nav-icon">👷</span>Workers</a><a href="{{url_for('contacts')}}" class="{{'active' if request.endpoint=='contacts' else ''}}"><span class="nav-icon">☎</span>Contacts</a><a href="{{url_for('ai_assistant')}}" class="{{'active' if request.endpoint=='ai_assistant' else ''}}"><span class="nav-icon">🤖</span>AI Assistant</a><a href="{{url_for('help_support')}}" class="{{'active' if request.endpoint=='help_support' else ''}}"><span class="nav-icon">?</span>Help & Support</a><a href="{{url_for('admin_profile')}}" class="{{'active' if request.endpoint=='admin_profile' else ''}}"><span class="nav-icon">⚙</span>Admin Profile</a><a href="{{url_for('admin_password_resets')}}" class="{{'active' if request.endpoint=='admin_password_resets' else ''}}"><span class="nav-icon">🔐</span>Password Requests</a>
{% elif session.get('role')=='worker' %}<a href="{{url_for('worker_dashboard')}}" class="{{'active' if request.endpoint=='worker_dashboard' else ''}}"><span class="nav-icon">▦</span>Assigned work</a><a href="{{url_for('notifications_page')}}"><span class="nav-icon">♢</span>Notifications</a><a href="{{url_for('contacts')}}"><span class="nav-icon">☎</span>Contacts</a><a href="{{url_for('ai_assistant')}}"><span class="nav-icon">🤖</span>AI Assistant</a><a href="{{url_for('help_support')}}"><span class="nav-icon">?</span>Help & Support</a>{% else %}<a href="{{url_for('student_dashboard')}}" class="{{'active' if request.endpoint=='student_dashboard' else ''}}"><span class="nav-icon">⌂</span>Dashboard</a><a href="{{url_for('new_complaint')}}" class="{{'active' if request.endpoint=='new_complaint' else ''}}"><span class="nav-icon">＋</span>New complaint</a><a href="{{url_for('contacts')}}" class="{{'active' if request.endpoint=='contacts' else ''}}"><span class="nav-icon">☎</span>Contacts</a><a href="{{url_for('ai_assistant')}}" class="{{'active' if request.endpoint=='ai_assistant' else ''}}"><span class="nav-icon">🤖</span>AI Assistant</a><a href="{{url_for('help_support')}}" class="{{'active' if request.endpoint=='help_support' else ''}}"><span class="nav-icon">?</span>Help & Support</a><a href="{{url_for('student_profile')}}" class="{{'active' if request.endpoint=='student_profile' else ''}}"><span class="nav-icon">◉</span>My Profile</a><a href="{{url_for('notices_page')}}"><span class="nav-icon">▤</span>Campus Notices</a><a href="{{url_for('notifications_page')}}"><span class="nav-icon">♢</span>Notifications</a><a href="{{url_for('lost_found_page')}}"><span class="nav-icon">⌕</span>Lost & Found</a><a href="{{url_for('campus_map')}}"><span class="nav-icon">⌖</span>Campus Map</a>{% endif %}{% if session.get('role')=='admin' %}<a href="{{url_for('notices_page')}}"><span class="nav-icon">▤</span>Campus Notices</a><a href="{{url_for('notifications_page')}}"><span class="nav-icon">♢</span>Notifications</a><a href="{{url_for('lost_found_page')}}"><span class="nav-icon">⌕</span>Lost & Found</a><a href="{{url_for('campus_map')}}"><span class="nav-icon">⌖</span>Campus Map</a>{% endif %}</nav><div class="side-bottom"><div class="user-mini"><div class="avatar">{{(session.get('name') or 'U')[0]|upper}}</div><div><strong>{{session.get('name')}}</strong><span>{{session.get('role')|capitalize}}</span></div></div><a class="logout" href="{{url_for('logout')}}">Sign out →</a></div></aside><main class="main"><header class="topbar"><div style="display:flex;align-items:center;gap:10px"><button class="mobile-menu" onclick="document.getElementById('sidebar').classList.toggle('open')">☰</button><span class="eyebrow">{{eyebrow or 'CAMPUS CARE 360'}}</span></div><div class="top-right"><span class="chip" style="background:#f5e4df;color:#8f202a">● {{'Admin mode' if session.get('role')=='admin' else ('Worker mode' if session.get('role')=='worker' else 'Student mode')}}</span></div></header>{% with messages=get_flashed_messages(with_categories=true) %}{% for category,message in messages %}<div class="flash {{category}}">{{message}}</div>{% endfor %}{% endwith %}<section class="content">{{body|safe}}</section></main></div>{% else %}<div class="auth-wrap">{% with messages=get_flashed_messages(with_categories=true) %}{% for category,message in messages %}<div class="flash {{category}}" style="position:fixed;top:20px;left:50%;transform:translateX(-50%);z-index:50;width:min(450px,90%)">{{message}}</div>{% endfor %}{% endwith %}{{body|safe}}</div>{% endif %}
<script>
(function(){
  function addPasswordToggles(){
    document.querySelectorAll('input[type="password"]:not([data-toggle-ready])').forEach(function(input){
      input.setAttribute('data-toggle-ready','true');
      var wrap=document.createElement('div');
      wrap.className='password-wrap';
      input.parentNode.insertBefore(wrap,input);
      wrap.appendChild(input);
      var button=document.createElement('button');
      button.type='button';
      button.className='password-toggle';
      button.textContent='Show';
      button.setAttribute('aria-label','Show password');
      button.setAttribute('aria-pressed','false');
      button.addEventListener('click',function(){
        var showing=input.type==='password';
        input.type=showing?'text':'password';
        button.textContent=showing?'Hide':'Show';
        button.setAttribute('aria-label',showing?'Hide password':'Show password');
        button.setAttribute('aria-pressed',showing?'true':'false');
        input.focus();
      });
      wrap.appendChild(button);
    });
  }
  if(document.readyState==='loading'){document.addEventListener('DOMContentLoaded',addPasswordToggles);}else{addPasswordToggles();}
})();
</script>
</body></html>'''

def render_page(body,title="Campus Care 360",eyebrow="CAMPUS CARE 360"):
    return render_template_string(SHELL,body=body,title=title,eyebrow=eyebrow)

def require(role):
    return session.get("role")==role

def ai_reply(message):
    m = (message or "").lower().strip()
    if not m:
        return "Hi! I’m Campus Care 360 AI. Ask about complaints, campus notices, lost & found, the campus map, login, or support."
    if any(x in m for x in ["lost", "found item", "lost and found"]):
        return "Open Lost & Found from the menu. You can report an item you lost or found, browse open listings, and send a claim request without exposing another student's contact details."
    if any(x in m for x in ["notice", "announcement", "holiday", "event"]):
        con = get_db()
        rows = con.execute("SELECT title,body,category,created_at FROM notices WHERE active=1 ORDER BY created_at DESC LIMIT 5").fetchall()
        con.close()
        if rows:
            return "Recent campus notices:\n" + "\n".join([f"• {r['title']} ({r['category']}): {r['body'][:180]}" for r in rows])
        return "There are no campus notices published yet. Open Campus Notices to check again later."
    if any(x in m for x in ["map", "where is", "location of", "find the library", "find hostel"]):
        return "Open Campus Map in the menu to browse campus buildings and service locations. If a location is missing, ask an administrator to add it."
    if any(x in m for x in ["my complaint", "my complaints", "complaint status", "track", "status"]):
        if session.get('role') == 'student':
            con = get_db()
            rows = con.execute("SELECT complaint_id,title,status,updated_at FROM complaints WHERE user_id=? ORDER BY updated_at DESC LIMIT 5", (session.get('user_id'),)).fetchall()
            con.close()
            if rows:
                return "Your recent complaint updates:\n" + "\n".join([f"• {r['complaint_id']} — {r['title']}: {r['status']} (updated {r['updated_at']})" for r in rows])
            return "You have not submitted a complaint yet. Use New Complaint to report an issue."
        return "Open your dashboard to view complaint status. Students can track their own reports; workers can view assigned tasks."
    if any(x in m for x in ["not working", "app problem", "error", "bug", "broken", "website"]):
        return "Try refreshing the page, checking your internet connection, and signing in again. If an upload fails, use PNG/JPG and a smaller image. If the problem remains, open Help & Support and share the page, steps, and any error message."
    if "login" in m or "password" in m or "forgot" in m:
        return "Use the campus email and password issued by your administrator. Choose Forgot Password to request a reset; the password changes only after admin approval. Never share your password in chat."
    if "complaint" in m and any(x in m for x in ["register", "submit", "new", "report"]):
        return "Open New Complaint, add a title, category, location, priority and description, then submit. Keep the complaint reference ID for follow-up."
    if "contact" in m or "phone" in m or "number" in m:
        return "Open Contacts or Emergency & SOS from the menu to find campus support numbers. For immediate danger in India, call 112. Confirm campus contact numbers with the administration."
    if "worker" in m or "staff" in m:
        return "Workers sign in through Worker Portal and see jobs assigned by the administrator. The admin creates worker login credentials and assigns complaints."
    if "student" in m and "admin" in m:
        return "Only administrators can create student accounts. Admins can manage accounts from the Students section."
    if "emergency" in m or "sos" in m or "ambulance" in m or "fire" in m:
        return "Open Emergency & SOS for campus safety information. If there is immediate danger in India, call 112. The app does not silently send an SOS alert."
    if "timetable" in m or "attendance" in m or "class" in m:
        return "Campus Care 360's current release focuses on support workflows. Ask the administrator to publish verified timetable and attendance data before relying on these details."
    return "I can help with complaints, notices, lost & found, campus locations, account access, workers and troubleshooting. Try asking: ‘Show my recent complaints’, ‘What are the latest notices?’, or ‘How do I report a lost item?’"

LOGIN = r'''<div class="auth-card"><div class="auth-brand"><img class="brand-logo-img" src="{{ url_for('static', filename='campus-care-360-logo.png') }}" alt="Campus Care 360 logo"><div><h1>Campus Care 360</h1><span>Hostel & campus support hub</span></div></div><h2>Welcome back.</h2><p>Report an issue, track progress and connect with campus support.</p><form method="POST"><div><label>Email</label><input type="email" name="email" placeholder="yourname@campus.360" required></div><div><label>Password</label><input type="password" name="password" placeholder="Enter your password" required></div><button class="btn btn-primary" style="justify-content:center" type="submit">Sign in →</button></form><div class="auth-link"><a href="{{url_for('forgot_password')}}">Forgot password?</a></div><div class="auth-link">Staff member? <a href="{{url_for('worker_login')}}">Open Worker Portal</a></div><div class="demo">Accounts are issued by Campus Care 360 administration only. Use your assigned email and password.</div></div>'''

REGISTER = r'''<div class="auth-card"><div class="auth-brand"><img class="brand-logo-img" src="{{ url_for('static', filename='campus-care-360-logo.png') }}" alt="Campus Care 360 logo"><div><h1>Campus Care 360</h1><span>Create your student support account</span></div></div><h2>Join Campus Care 360.</h2><p>Your registration details will be visible to authorized admins for support management.</p><form method="POST"><div><label>Full name *</label><input name="name" required placeholder="Your full name"></div><div><label>Email *</label><input type="email" name="email" required placeholder="yourname@campus.360"></div><div><label>Mobile number *</label><input name="mobile" required pattern="[0-9]{10}" maxlength="10" placeholder="10-digit mobile number"></div><div class="form-grid"><div><label>Student ID</label><input name="student_id" placeholder="CC2026002"></div><div><label>Hostel / Room</label><input name="room" placeholder="B-204"></div><div><label>Department</label><input name="department" placeholder="CSE"></div><div><label>Year</label><select name="year"><option value="">Select year</option><option>1st Year</option><option>2nd Year</option><option>3rd Year</option><option>4th Year</option></select></div></div><div><label>Password *</label><input type="password" name="password" required placeholder="Create a password"></div><button class="btn btn-blue" style="justify-content:center" type="submit">Create account →</button></form><div class="auth-link">Already registered? <a href="{{url_for('login')}}">Sign in</a></div></div>'''

@app.route('/')
def home():
    # Require an authenticated account before opening any dashboard.
    if not session.get('user_id') and not session.get('worker_id'):
        return redirect(url_for('login'))

    role = session.get('role')
    if role == 'worker':
        return redirect(url_for('worker_dashboard'))
    if role == 'admin':
        return redirect(url_for('admin_dashboard'))
    if role == 'student':
        return redirect(url_for('student_dashboard'))

    # Unknown or incomplete session: clear it and require sign-in again.
    session.clear()
    return redirect(url_for('login'))

@app.route('/login',methods=['GET','POST'])
def login():
    if request.method=='POST':
        email=request.form.get('email','').strip().lower(); password=request.form.get('password','')
        if not valid_campus_email(email):
            flash("Use an email ending with @campus.360.","error"); return render_page(render_template_string(LOGIN),'Login')
        con=get_db(); u=con.execute('SELECT * FROM users WHERE email=?',(email,)).fetchone(); con.close()
        if u and check_password_hash(u['password'],password):
            session.clear(); session.update(user_id=u['id'],name=u['name'],role=u['role']); return redirect(url_for('home'))
        flash('Invalid email or password.','error')
    return render_page(render_template_string(LOGIN),'Login')

@app.route('/forgot-password', methods=['GET', 'POST'])
def forgot_password():
    if request.method == 'POST':
        email = request.form.get('email', '').strip().lower()
        new_password = request.form.get('new_password', '')
        confirm_password = request.form.get('confirm_password', '')
        if not valid_campus_email(email):
            flash('Enter your registered campus email ending in @campus.360.', 'error')
            return redirect(url_for('forgot_password'))
        if len(new_password) < 8:
            flash('Choose a password with at least 8 characters.', 'error')
            return redirect(url_for('forgot_password'))
        if new_password != confirm_password:
            flash('The passwords do not match.', 'error')
            return redirect(url_for('forgot_password'))
        con = get_db()
        user = con.execute("SELECT id,email FROM users WHERE lower(email)=? AND role='student'", (email,)).fetchone()
        account_type = 'student'
        account_id = user['id'] if user else None
        if not user:
            worker = con.execute('SELECT id,email FROM workers WHERE lower(email)=? AND active=1', (email,)).fetchone()
            if worker:
                account_type = 'worker'
                account_id = worker['id']
        if account_id is None:
            con.close()
            # Avoid revealing whether an email exists.
            flash('If the account exists, your reset request can be reviewed by the administrator.', 'success')
            return redirect(url_for('login'))
        pending = con.execute("SELECT id FROM password_reset_requests WHERE account_type=? AND account_id=? AND status='Pending'", (account_type, account_id)).fetchone()
        if pending:
            con.close()
            flash('A password reset request is already awaiting admin approval.', 'success')
            return redirect(url_for('login'))
        con.execute('INSERT INTO password_reset_requests(account_type,account_id,email,pending_password_hash) VALUES(?,?,?,?)',
                     (account_type, account_id, email, generate_password_hash(new_password)))
        con.commit()
        con.close()
        flash('Your password reset request has been sent to the administrator for approval. Your password will not change until it is approved.', 'success')
        return redirect(url_for('login'))
    body = render_template_string(r'''<div class="auth-card"><div class="auth-brand"><img class="brand-logo-img" src="{{ url_for('static', filename='campus-care-360-logo.png') }}" alt="Campus Care 360 logo"><div><h1>Campus Care 360</h1><span>Secure account recovery</span></div></div><h2>Reset password.</h2><p>Submit a reset request. An administrator must approve it before your password changes.</p><form method="POST"><div><label>Registered campus email</label><input type="email" name="email" required placeholder="you@campus.360" pattern="[^@\s]+@campus\.360"></div><div><label>New password</label><input type="password" name="new_password" minlength="8" required autocomplete="new-password" placeholder="At least 8 characters"></div><div><label>Confirm new password</label><input type="password" name="confirm_password" minlength="8" required autocomplete="new-password" placeholder="Re-enter new password"></div><button class="btn btn-primary" style="justify-content:center" type="submit">Send approval request →</button></form><div class="demo">Your new password remains pending until the Campus Care 360 administrator verifies and approves this request.</div><div class="auth-link"><a href="{{url_for('login')}}">← Back to login</a></div></div>''')
    return render_page(body, 'Forgot Password', 'ACCOUNT RECOVERY')

@app.route('/admin/password-resets', methods=['GET', 'POST'])
def admin_password_resets():
    if not require('admin'):
        return redirect(url_for('login'))
    con = get_db()
    if request.method == 'POST':
        request_id = request.form.get('request_id', '').strip()
        action = request.form.get('action', '').strip()
        reset = con.execute("SELECT * FROM password_reset_requests WHERE id=? AND status='Pending'", (request_id,)).fetchone() if request_id.isdigit() else None
        if not reset:
            flash('That pending request could not be found.', 'error')
        elif action == 'approve':
            if reset['account_type'] == 'student':
                exists = con.execute("SELECT id FROM users WHERE id=? AND role='student'", (reset['account_id'],)).fetchone()
                if exists:
                    con.execute('UPDATE users SET password=? WHERE id=?', (reset['pending_password_hash'], reset['account_id']))
                else:
                    exists = None
            else:
                exists = con.execute('SELECT id FROM workers WHERE id=? AND active=1', (reset['account_id'],)).fetchone()
                if exists:
                    con.execute('UPDATE workers SET password=? WHERE id=?', (reset['pending_password_hash'], reset['account_id']))
            if exists:
                con.execute("UPDATE password_reset_requests SET status='Approved', reviewed_at=CURRENT_TIMESTAMP WHERE id=?", (reset['id'],))
                con.commit()
                if reset['account_type'] == 'student': create_notification('user', reset['account_id'], 'Password reset approved', 'Your password reset was approved. You can now sign in with the new password.', url_for('login'))
                else: create_notification('worker', reset['account_id'], 'Password reset approved', 'Your password reset was approved. You can now sign in with the new password.', url_for('worker_login'))
                flash('Password reset approved for ' + reset['email'] + '. The new password is now active.', 'success')
            else:
                con.execute("UPDATE password_reset_requests SET status='Rejected', reviewed_at=CURRENT_TIMESTAMP WHERE id=?", (reset['id'],))
                con.commit()
                flash('Account no longer exists or is inactive. Request rejected.', 'error')
        elif action == 'reject':
            con.execute("UPDATE password_reset_requests SET status='Rejected', reviewed_at=CURRENT_TIMESTAMP WHERE id=?", (reset['id'],))
            con.commit()
            flash('Password reset request rejected.', 'success')
        else:
            flash('Unknown action.', 'error')
    requests = con.execute('''SELECT * FROM password_reset_requests ORDER BY CASE status WHEN 'Pending' THEN 0 ELSE 1 END, created_at DESC''').fetchall()
    con.close()
    body = render_template_string(r'''<div class="page-head"><div><div class="eyebrow">ACCOUNT SECURITY</div><h1>Password reset requests</h1><p>Verify the requester’s identity before approving a password change.</p></div></div><div class="admin-list">{% for r in requests %}<div class="admin-card"><div class="admin-top"><div><div class="admin-meta">REQUEST #{{r['id']}} · {{r['account_type']|title}} · {{r['created_at']}}</div><div class="admin-title">{{r['email']}}</div><div class="admin-meta">Requested status: {{r['status']}}</div></div><span class="chip {{'progress' if r['status']=='Pending' else ('resolved' if r['status']=='Approved' else 'closed')}}">{{r['status']}}</span></div>{% if r['status']=='Pending' %}<p class="admin-desc">The user’s proposed new password is stored as a hash and cannot be viewed here. Confirm their identity through a trusted channel before approving.</p><form method="POST" style="display:flex;gap:10px;flex-wrap:wrap"><input type="hidden" name="request_id" value="{{r['id']}}"><button class="btn btn-primary" name="action" value="approve" type="submit" onclick="return confirm('Verify this user first. Approve this password reset?')">Approve reset</button><button class="btn btn-light" name="action" value="reject" type="submit" onclick="return confirm('Reject this password reset request?')">Reject</button></form>{% else %}<div class="admin-meta">Reviewed: {{r['reviewed_at'] or '—'}}</div>{% endif %}</div>{% else %}<div class="empty">No password reset requests yet.</div>{% endfor %}</div>''', requests=requests)
    return render_page(body, 'Password Reset Requests', 'ACCOUNT SECURITY')

@app.route('/register', methods=['GET', 'POST'])
def register():
    # Student accounts are created by administrators only.
    flash('Student accounts are issued by Campus Care 360 administration. Please contact your administrator.', 'error')
    return redirect(url_for('login'))

@app.route('/worker/login', methods=['GET','POST'])
def worker_login():
    if request.method == 'POST':
        email = request.form.get('email','').strip().lower()
        password = request.form.get('password','')
        con = get_db()
        worker = con.execute('SELECT * FROM workers WHERE lower(email)=? AND active=1', (email,)).fetchone()
        con.close()
        if worker and worker['password'] and check_password_hash(worker['password'], password):
            session.clear()
            session.update(worker_id=worker['id'], name=worker['name'], role='worker')
            return redirect(url_for('worker_dashboard'))
        flash('Invalid worker email or password. Ask the administrator for your login.', 'error')
    body = render_template_string(r'''<div class="auth-card"><div class="auth-brand"><img class="brand-logo-img" src="{{ url_for('static', filename='campus-care-360-logo.png') }}" alt="Campus Care 360 logo"><div><h1>Campus Care 360</h1><span>Worker operations portal</span></div></div><h2>Worker sign in.</h2><p>View assigned jobs and update their progress.</p><form method="POST"><div><label>Work email</label><input type="email" name="email" required placeholder="worker@campus.360"></div><div><label>Password</label><input type="password" name="password" required placeholder="Password issued by admin"></div><button class="btn btn-primary" style="justify-content:center" type="submit">Enter worker portal →</button></form><div class="auth-link"><a href="{{url_for('login')}}">← Student / Admin login</a></div><div class="demo">Worker accounts can only be created by an authorized administrator.</div></div>''')
    return render_page(body, 'Worker Login', 'WORKER PORTAL')

@app.route('/worker')
def worker_dashboard():
    if not require('worker') or not session.get('worker_id'):
        return redirect(url_for('worker_login'))
    con = get_db()
    jobs = con.execute("SELECT c.*,u.name student_name,u.room FROM complaints c JOIN users u ON u.id=c.user_id WHERE c.assigned_worker_id=? ORDER BY CASE c.priority WHEN 'Emergency' THEN 0 WHEN 'High' THEN 1 ELSE 2 END, c.created_at DESC", (session['worker_id'],)).fetchall()
    con.close()
    body = render_template_string(r'''<div class="page-head"><div><div class="eyebrow">WORKER OPERATIONS</div><h1>Welcome, {{session.get('name','Worker').split()[0]}}.</h1><p>Your assigned campus tasks, prioritized in one place.</p></div><span class="chip emergency">● WORKER PORTAL</span></div><div class="hero"><div class="hero-content"><div class="hero-kicker">CAMPUS CARE / FIELD TEAM</div><h2>Every task. Clear ownership.</h2><p>Review the issue, visit the reported location, and keep students updated as work progresses.</p></div></div><div class="emergency-actions"><a class="emergency-tile sos" href="tel:112"><span class="tile-icon">🆘</span><span><strong>SOS · Call 112</strong><small>Use in immediate danger or urgent medical/fire situations.</small></span><span class="tile-arrow">↗</span></a><a class="emergency-tile" href="{{url_for('emergency_help')}}"><span class="tile-icon">✚</span><span><strong>Emergency Help Centre</strong><small>View emergency guidance and support numbers.</small></span><span class="tile-arrow">›</span></a></div><div class="section-title"><h3>Your assigned jobs</h3><span>{{jobs|length}} tasks</span></div><div class="admin-list">{% for j in jobs %}<div class="admin-card"><div class="admin-top"><div><div class="admin-meta">{{j['complaint_id']}} · {{j['category']}} · {{j['location']}}</div><div class="admin-title">{{j['title']}}</div><div class="admin-meta">Reported by {{j['student_name']}} · Room {{j['room'] or '—'}}</div></div><div class="admin-tags"><span class="chip {{'emergency' if j['priority']=='Emergency' else 'assigned'}}">{{j['priority']}}</span><span class="chip progress">{{j['status']}}</span></div></div><p class="admin-desc">{{j['description']}}</p><form class="admin-form" method="POST" action="{{url_for('worker_update_job', id=j['id'])}}"><div><label>Update status</label><select name="status">{% for st in ['Assigned','In Progress','Resolved'] %}<option {{'selected' if j['status']==st else ''}}>{{st}}</option>{% endfor %}</select></div><div><label>Update note</label><input name="remark" placeholder="Work done / next step" value="{{j['admin_remark']}}"></div><button class="btn btn-primary" type="submit">Save update</button></form></div>{% else %}<div class="empty">No jobs assigned yet. Ask the administrator to assign a complaint to your worker account.</div>{% endfor %}</div>''', jobs=jobs)
    return render_page(body, 'Worker Portal', 'WORKER OPERATIONS')

@app.route('/worker/job/<int:id>', methods=['POST'])
def worker_update_job(id):
    if not require('worker') or not session.get('worker_id'):
        return redirect(url_for('worker_login'))
    status = request.form.get('status','In Progress')
    if status not in ('Assigned','In Progress','Resolved'):
        status = 'In Progress'
    remark = request.form.get('remark','').strip()
    con = get_db()
    job = con.execute('SELECT user_id,complaint_id,title FROM complaints WHERE id=? AND assigned_worker_id=?',(id,session['worker_id'])).fetchone()
    cur = con.execute("UPDATE complaints SET status=?,admin_remark=?,updated_at=CURRENT_TIMESTAMP WHERE id=? AND assigned_worker_id=?", (status, remark, id, session['worker_id']))
    con.commit(); con.close()
    if job and cur.rowcount: create_notification('user', job['user_id'], 'Worker updated your complaint', f"{job['complaint_id']} · {job['title']} is now {status}.", url_for('complaint_status', complaint_id=job['complaint_id']))
    flash('Task updated.' if cur.rowcount else 'That task is not assigned to your account.', 'success' if cur.rowcount else 'error')
    return redirect(url_for('worker_dashboard'))

@app.route('/logout')
def logout(): session.clear(); return redirect(url_for('login'))

@app.route('/student')
def student_dashboard():
    if not require('student'): return redirect(url_for('login'))
    con=get_db(); cs=con.execute('SELECT * FROM complaints WHERE user_id=? ORDER BY created_at DESC',(session['user_id'],)).fetchall(); notices=con.execute('SELECT * FROM notices WHERE active=1 ORDER BY created_at DESC LIMIT 4').fetchall(); unread=con.execute("SELECT COUNT(*) n FROM notifications WHERE account_type='user' AND account_id=? AND is_read=0",(session['user_id'],)).fetchone()['n']; con.close(); counts={s:sum(1 for c in cs if c['status']==s) for s in STATUSES}
    body=render_template_string(r'''<div class="page-head"><div><div class="eyebrow">STUDENT SPACE</div><h1>Good day, {{session.get('name','Student').split()[0]}} 👋</h1><p>Everything you report, in one calm place.</p></div><a class="btn btn-primary" href="{{url_for('new_complaint')}}">＋ Report an issue</a></div><div class="hero"><div class="hero-content"><div class="hero-kicker">CAMPUS CARE</div><h2>Make your campus better,<br>one report at a time.</h2><p>Report hostel and campus issues, attach evidence and follow every update without chasing anyone.</p><a class="hero-link" href="{{url_for('new_complaint')}}">Start a complaint →</a></div><div class="orbit"></div></div><div class="emergency-actions"><a class="emergency-tile sos" href="tel:112"><span class="tile-icon">🆘</span><span><strong> SOS · Call Emergency Services</strong><small>Tap to call India’s emergency number 112 if you are in immediate danger.</small></span><span class="tile-arrow">↗</span></a><a class="emergency-tile" href="{{url_for('emergency_help')}}"><span class="tile-icon">✚</span><span><strong>Emergency Help Centre</strong><small>Fire, ambulance, security and campus support contacts.</small></span><span class="tile-arrow">›</span></a></div><div class="section-title"><h3>Your campus today</h3><span>{{unread}} unread notifications</span></div><div class="help-quicklinks"><a class="help-quicklink" href="{{url_for('notices_page')}}"><span class="help-quick-icon">▤</span><span><b>Campus notices</b><small>Latest announcements</small></span></a><a class="help-quicklink" href="{{url_for('notifications_page')}}"><span class="help-quick-icon">♢</span><span><b>Notifications · {{unread}}</b><small>Complaint and campus updates</small></span></a><a class="help-quicklink" href="{{url_for('lost_found_page')}}"><span class="help-quick-icon">⌕</span><span><b>Lost & Found</b><small>Find or report an item</small></span></a><a class="help-quicklink" href="{{url_for('campus_map')}}"><span class="help-quick-icon">⌖</span><span><b>Campus Map</b><small>Find campus services</small></span></a></div><div class="section-title"><h3>Recent campus notices</h3><span><a href="{{url_for('notices_page')}}">View all →</a></span></div><div class="complaints">{% for n in notices %}<div class="complaint-row"><div><div class="row-meta">{{n['category']}} · {{n['created_at']}}</div><div class="row-title">{{n['title']}}</div><div class="row-sub">{{n['body']}}</div></div><span class="chip {{'emergency' if n['priority']=='Urgent' else 'submitted'}}">{{n['priority']}}</span></div>{% else %}<div class="empty">No campus notices have been published yet.</div>{% endfor %}</div><div class="section-title"><h3>Your overview</h3><span>{{cs|length}} total reports</span></div><div class="stats"><div class="stat"><div class="stat-top"><span>Submitted</span><div class="stat-icon">↗</div></div><strong>{{counts['Submitted']}}</strong></div><div class="stat"><div class="stat-top"><span>Being handled</span><div class="stat-icon">◌</div></div><strong>{{counts['Assigned']+counts['In Progress']}}</strong></div><div class="stat"><div class="stat-top"><span>Resolved</span><div class="stat-icon">✓</div></div><strong>{{counts['Resolved']}}</strong></div><div class="stat"><div class="stat-top"><span>Closed</span><div class="stat-icon">▣</div></div><strong>{{counts['Closed']}}</strong></div></div><div class="section-title"><h3>Recent complaints</h3><span>Open a report for details</span></div><div class="complaints">{% for c in cs %}<a class="complaint-row" href="{{url_for('complaint_status',complaint_id=c['complaint_id'])}}"><div><div class="row-meta">{{c['category']|upper}} · {{c['location']}} · {{c['complaint_id']}}</div><div class="row-title">{{c['title']}}</div><div class="row-sub">{{c['created_at']}}</div></div><div class="row-right"><span class="chip {{ {'Submitted':'submitted','Assigned':'assigned','In Progress':'progress','Resolved':'resolved','Closed':'closed'}[c['status']] }}">{{c['status']}}</span><span class="arrow">›</span></div></a>{% else %}<div class="empty">No complaints yet. Your first report can start here.</div>{% endfor %}</div><a class="floating-ai" title="AI Assistant" href="{{url_for('ai_assistant')}}">🤖</a>''',cs=cs,counts=counts,notices=notices,unread=unread)
    return render_page(body,'Student Dashboard','STUDENT SPACE')

@app.route('/complaint/new',methods=['GET','POST'])
def new_complaint():
    if not require('student'): return redirect(url_for('login'))
    if request.method=='POST':
        title=request.form.get('title','').strip(); desc=request.form.get('description','').strip(); category=request.form.get('category',''); location=request.form.get('location','').strip(); priority=request.form.get('priority','Medium'); f=request.files.get('image'); image=''
        if not title or not desc or not category or not location: flash('Please complete all required fields.','error'); return redirect(url_for('new_complaint'))
        if f and f.filename:
            if not allowed_file(f.filename): flash('Please upload PNG, JPG, JPEG, GIF or WEBP.','error'); return redirect(url_for('new_complaint'))
            image=secure_filename(uuid.uuid4().hex[:10]+'_'+f.filename); f.save(os.path.join(UPLOAD_FOLDER,image))
        cid='CC-'+uuid.uuid4().hex[:8].upper(); con=get_db(); con.execute('INSERT INTO complaints(complaint_id,user_id,category,title,description,location,priority,image) VALUES(?,?,?,?,?,?,?,?)',(cid,session['user_id'],category,title,desc,location,priority,image)); con.commit(); con.close(); create_notification('user', session['user_id'], 'Complaint submitted', f'Your report {cid} was received.', url_for('complaint_status', complaint_id=cid)); flash(f'Complaint {cid} submitted successfully.','success'); return redirect(url_for('complaint_status',complaint_id=cid))
    body=render_template_string(r'''<div class="page-head"><div><div class="eyebrow">NEW REPORT</div><h1>Tell us what needs attention</h1><p>A clear report helps the campus team solve it faster.</p></div></div><div class="card"><form method="POST" enctype="multipart/form-data"><div class="form-grid"><div class="full"><label>What is the issue? *</label><input name="title" required placeholder="Example: Water leaking near room B-204"></div><div><label>Category *</label><select name="category" required><option value="">Choose a category</option>{% for c in categories %}<option>{{c}}</option>{% endfor %}</select></div><div><label>Location *</label><input name="location" required placeholder="Block B, Room 204"></div><div><label>Priority</label><select name="priority">{% for p in priorities %}<option>{{p}}</option>{% endfor %}</select></div><div><label>Photo evidence</label><div class="upload">Attach an image if useful<input type="file" name="image" accept="image/*"></div></div><div class="full"><label>Description *</label><textarea name="description" required placeholder="Explain what happened, when you noticed it and anything that may help the staff."></textarea></div></div><div class="form-actions"><a class="btn btn-light" href="{{url_for('student_dashboard')}}">Cancel</a><button class="btn btn-primary" type="submit">Submit complaint →</button></div></form></div>''')
    return render_page(body,'New Complaint','NEW REPORT')

@app.route('/complaint/<complaint_id>')
def complaint_status(complaint_id):
    if not session.get('user_id'): return redirect(url_for('login'))
    con=get_db(); c=con.execute('''SELECT c.*,u.name,u.email,u.room,u.mobile,w.name worker_name,w.department worker_dept,w.mobile worker_mobile FROM complaints c JOIN users u ON u.id=c.user_id LEFT JOIN workers w ON w.id=c.assigned_worker_id WHERE c.complaint_id=?''',(complaint_id,)).fetchone(); feedback=con.execute('SELECT * FROM feedback WHERE complaint_id=? ORDER BY id DESC LIMIT 1',(c['id'],)).fetchone() if c else None; con.close()
    if not c: flash('Complaint not found.','error'); return redirect(url_for('home'))
    if session.get('role')=='student' and c['user_id']!=session['user_id']: flash('You cannot view that complaint.','error'); return redirect(url_for('student_dashboard'))
    idx=STATUSES.index(c['status']) if c['status'] in STATUSES else 0
    body=render_template_string(r'''<div class="page-head"><div><div class="eyebrow">COMPLAINT STATUS</div><h1>{{c['complaint_id']}}</h1><p>Track the report from submission to closure.</p></div><a class="btn btn-light" href="{{url_for('home')}}">← Back</a></div><div class="detail-grid"><div class="card"><div class="detail-title"><div><h1>{{c['title']}}</h1><div class="mini-meta">{{c['category']}} · {{c['location']}}</div></div><span class="chip {{ {'Submitted':'submitted','Assigned':'assigned','In Progress':'progress','Resolved':'resolved','Closed':'closed'}[c['status']] }}">{{c['status']}}</span></div><div class="detail-block"><h4>Progress</h4><div class="timeline">{% for s in statuses %}<div class="step {{'done' if loop.index0<idx else ('current' if loop.index0==idx else '')}}"><div class="dot">{{'✓' if loop.index0<idx else loop.index}}</div><div><b>{{s}}</b><span>{{'Completed' if loop.index0<idx else ('Current status' if loop.index0==idx else 'Waiting for next step')}}</span></div></div>{% endfor %}</div></div><div class="detail-block"><h4>Report details</h4><p>{{c['description']}}</p></div>{% if c['image'] %}<div class="detail-block"><h4>Photo evidence</h4><img class="photo" src="{{url_for('uploaded_file',filename=c['image'])}}"></div>{% endif %}{% if c['admin_remark'] %}<div class="detail-block"><h4>Admin remark</h4><p>{{c['admin_remark']}}</p></div>{% endif %}{% if c['status']=='Resolved' and session.get('role')=='student' and not feedback %}<div class="feedback"><h4>How did we do?</h4><form method="POST" action="{{url_for('submit_feedback',complaint_id=c['complaint_id'])}}"><div style="display:flex;gap:7px;font-size:25px"><label><input type="radio" name="rating" value="1" required>★</label><label><input type="radio" name="rating" value="2">★★</label><label><input type="radio" name="rating" value="3">★★★</label><label><input type="radio" name="rating" value="4">★★★★</label><label><input type="radio" name="rating" value="5">★★★★★</label></div><textarea name="comment" style="min-height:80px;margin-top:10px" placeholder="Optional comment"></textarea><button class="btn btn-primary" style="margin-top:10px" type="submit">Send feedback</button></form></div>{% elif feedback %}<div class="feedback"><h4>Feedback received</h4><div style="font-size:22px">{{'★'*feedback['rating']}}{{'☆'*(5-feedback['rating'])}}</div><p>{{feedback['comment']}}</p></div>{% endif %}</div><div style="display:grid;gap:18px;align-content:start"><div class="card"><h3 style="margin:0 0 17px">Assignment</h3><div class="info-list"><div class="info-item"><span>Priority</span><strong>{{c['priority']}}</strong></div><div class="info-item"><span>Assigned to</span><strong>{{c['worker_name'] or c['assigned_to'] or 'Not assigned yet'}}</strong></div>{% if c['worker_mobile'] %}<div class="info-item"><span>Worker mobile</span><strong>{{c['worker_mobile']}}</strong></div>{% endif %}<div class="info-item"><span>Student</span><strong>{{c['name']}}</strong></div><div class="info-item"><span>Room</span><strong>{{c['room'] or '—'}}</strong></div><div class="info-item"><span>Created</span><strong>{{c['created_at']}}</strong></div><div class="info-item"><span>Updated</span><strong>{{c['updated_at']}}</strong></div></div></div><div class="card"><h3 style="margin:0 0 8px">Need help?</h3><p class="muted">Keep <b>{{c['complaint_id']}}</b> when speaking with campus support.</p><a class="btn btn-light" href="{{url_for('ai_assistant')}}">Ask AI Assistant</a></div></div></div>''',c=c,idx=idx,statuses=STATUSES,feedback=feedback)
    return render_page(body,'Complaint Status','COMPLAINT STATUS')

@app.route('/complaint/<complaint_id>/feedback',methods=['POST'])
def submit_feedback(complaint_id):
    if not require('student'): return redirect(url_for('login'))
    try: rating=max(1,min(5,int(request.form.get('rating','5'))))
    except: rating=5
    con=get_db(); c=con.execute('SELECT * FROM complaints WHERE complaint_id=? AND user_id=?',(complaint_id,session['user_id'])).fetchone()
    if c: con.execute('INSERT INTO feedback(complaint_id,rating,comment) VALUES(?,?,?)',(c['id'],rating,request.form.get('comment','').strip())); con.execute("UPDATE complaints SET status='Closed',updated_at=CURRENT_TIMESTAMP WHERE id=?",(c['id'],)); con.commit()
    con.close(); flash('Thank you for your feedback. The complaint is now closed.','success'); return redirect(url_for('complaint_status',complaint_id=complaint_id))



@app.route('/admin')
def admin_dashboard():
    if not require('admin') or not session.get('user_id'):
        return redirect(url_for('login'))

    con = get_db()
    total = con.execute('SELECT COUNT(*) AS n FROM complaints').fetchone()['n']
    submitted = con.execute("SELECT COUNT(*) AS n FROM complaints WHERE status='Submitted'").fetchone()['n']
    active = con.execute("SELECT COUNT(*) AS n FROM complaints WHERE status IN ('Assigned','In Progress')").fetchone()['n']
    resolved = con.execute("SELECT COUNT(*) AS n FROM complaints WHERE status IN ('Resolved','Closed')").fetchone()['n']
    complaints = con.execute('''
        SELECT c.*, u.name AS student_name, u.room,
               w.name AS worker_name
        FROM complaints c
        JOIN users u ON u.id = c.user_id
        LEFT JOIN workers w ON w.id = c.assigned_worker_id
        ORDER BY CASE c.priority WHEN 'Emergency' THEN 0 WHEN 'High' THEN 1 ELSE 2 END,
                 c.created_at DESC
        LIMIT 100
    ''').fetchall()
    workers = con.execute('SELECT id, name, department, mobile FROM workers WHERE active=1 ORDER BY name').fetchall()
    con.close()

    body = render_template_string(r'''<div class="page-head">
      <div><div class="eyebrow">ADMINISTRATION</div><h1>Admin Dashboard</h1>
      <p>Review campus complaints, assign workers and track resolutions.</p></div>
      <span class="chip emergency">ADMIN PORTAL</span>
    </div>
    <div class="stats">
      <div class="stat"><div class="stat-top"><span>Total complaints</span><div class="stat-icon">▦</div></div><strong>{{total}}</strong></div>
      <div class="stat"><div class="stat-top"><span>Awaiting review</span><div class="stat-icon">↗</div></div><strong>{{submitted}}</strong></div>
      <div class="stat"><div class="stat-top"><span>Being handled</span><div class="stat-icon">◌</div></div><strong>{{active}}</strong></div>
      <div class="stat"><div class="stat-top"><span>Resolved / closed</span><div class="stat-icon">✓</div></div><strong>{{resolved}}</strong></div>
    </div>
    <div class="card" style="margin-bottom:24px;border-color:#e5d5ad;background:linear-gradient(135deg,#fffdf7,#fff9e9)">
      <div class="section-title"><h3>📊 Import accounts from Excel</h3><span>Admin only</span></div>
      <p class="muted">Download the Excel template to prepare student and worker account details in the correct format. The file-upload control has been removed from this dashboard.</p>
      <div style="display:flex;gap:12px;flex-wrap:wrap;align-items:center">
        <a class="btn btn-light" href="{{url_for('admin_import_template')}}">⬇ Download Excel template</a>
      </div>
      <p class="muted" style="margin-bottom:0">The workbook includes <b>Students</b> and <b>Workers</b> sheets with the required column headings.</p>
    </div>
    <div class="section-title"><h3>Complaint management</h3><span>{{complaints|length}} recent reports</span></div>
    <div class="admin-list">
    {% for c in complaints %}
      <article class="admin-card">
        <div class="admin-top"><div><div class="admin-meta">{{c['complaint_id']}} · {{c['category']}} · {{c['location']}}</div>
        <div class="admin-title">{{c['title']}}</div><div class="admin-meta">Reported by {{c['student_name']}} · Room {{c['room'] or '—'}} · {{c['created_at']}}</div></div>
        <div class="admin-tags"><span class="chip {{'emergency' if c['priority']=='Emergency' else ('progress' if c['priority']=='High' else 'submitted')}}">{{c['priority']}}</span>
        <span class="chip {{ {'Submitted':'submitted','Assigned':'assigned','In Progress':'progress','Resolved':'resolved','Closed':'closed'}.get(c['status'],'submitted') }}">{{c['status']}}</span></div></div>
        <p class="admin-desc">{{c['description']}}</p>
        <form class="admin-form" method="POST" action="{{url_for('update_complaint', id=c['id'])}}">
          <div><label>Status</label><select name="status">{% for st in statuses %}<option value="{{st}}" {{'selected' if c['status']==st else ''}}>{{st}}</option>{% endfor %}</select></div>
          <div><label>Assign worker</label><select name="worker_id"><option value="">Unassigned</option>{% for w in workers %}<option value="{{w['id']}}" {{'selected' if c['assigned_worker_id']==w['id'] else ''}}>{{w['name']}} · {{w['department']}}</option>{% endfor %}</select></div>
          <div><label>Admin remark</label><input name="remark" value="{{c['admin_remark'] or ''}}" placeholder="Next action or resolution note"></div>
          <button class="btn btn-primary" type="submit">Save update</button>
        </form>
      </article>
    {% else %}<div class="empty">No complaints have been submitted yet.</div>{% endfor %}
    </div>''', total=total, submitted=submitted, active=active, resolved=resolved,
       complaints=complaints, workers=workers, statuses=STATUSES)
    return render_page(body, 'Admin Dashboard', 'ADMINISTRATION')


@app.route('/admin/complaint/<int:id>',methods=['POST'])
def update_complaint(id):
    if not require('admin'): return redirect(url_for('login'))
    status=request.form.get('status','Submitted'); status=status if status in STATUSES else 'Submitted'; wid=request.form.get('worker_id','').strip(); remark=request.form.get('remark','').strip(); con=get_db(); worker=con.execute('SELECT * FROM workers WHERE id=?',(wid,)).fetchone() if wid.isdigit() else None
    row=con.execute('SELECT user_id,complaint_id,title FROM complaints WHERE id=?',(id,)).fetchone(); con.execute('UPDATE complaints SET status=?,assigned_worker_id=?,assigned_to=?,admin_remark=?,updated_at=CURRENT_TIMESTAMP WHERE id=?',(status,worker['id'] if worker else None,worker['name'] if worker else '',remark,id)); con.commit(); con.close();
    if row: create_notification('user', row['user_id'], 'Complaint update', f"{row['complaint_id']} · {row['title']} is now {status}.", url_for('complaint_status', complaint_id=row['complaint_id']))
    if worker: create_notification('worker', worker['id'], 'New assigned task', f"{row['complaint_id'] if row else 'Complaint'}: {row['title'] if row else ''}", url_for('worker_dashboard'))
    flash('Complaint updated successfully.','success'); return redirect(url_for('admin_dashboard'))

@app.route('/admin/import-template.xlsx')
def admin_import_template():
    if not require('admin'):
        return redirect(url_for('login'))
    try:
        from openpyxl import Workbook
    except ImportError:
        flash('Excel support is not installed. Add openpyxl to requirements.txt.', 'error')
        return redirect(url_for('admin_dashboard'))
    wb = Workbook()
    ws = wb.active
    ws.title = 'Students'
    ws.append(['name','email','password','mobile','room','department','year','student_id'])
    ws.append(['Example Student','student1@campus.360','ChangeMe123','9876543210','B-204','CSE','1st Year','CC2026001'])
    ww = wb.create_sheet('Workers')
    ww.append(['name','email','password','mobile','department','worker_code'])
    ww.append(['Example Worker','worker1@campus.360','ChangeMe123','9876543211','Electrical','WKEXAMPLE'])
    for sheet in wb.worksheets:
        sheet.freeze_panes='A2'
        sheet.auto_filter.ref=sheet.dimensions
        for col in sheet.columns:
            letter=col[0].column_letter
            sheet.column_dimensions[letter].width=min(max(max(len(str(c.value or '')) for c in col)+3,14),30)
    output=io.BytesIO(); wb.save(output); output.seek(0)
    return Response(output.getvalue(), mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition':'attachment; filename=Campus_Care_360_Import_Template.xlsx'})

@app.route('/admin/import-data', methods=['POST'])
def admin_import_data():
    if not require('admin') or not session.get('user_id'):
        return redirect(url_for('login'))
    uploaded=request.files.get('import_file')
    if not uploaded or not uploaded.filename:
        flash('Choose an Excel (.xlsx) or CSV file first.', 'error')
        return redirect(url_for('admin_dashboard'))
    filename=uploaded.filename.lower()
    try:
        if filename.endswith('.xlsx'):
            from openpyxl import load_workbook
            wb=load_workbook(uploaded, read_only=True, data_only=True)
            datasets={}
            for sheet_name in ('Students','Workers'):
                if sheet_name in wb.sheetnames:
                    rows=wb[sheet_name].iter_rows(values_only=True)
                    headers=next(rows, None)
                    if headers:
                        keys=[str(v or '').strip().lower().replace(' ','_') for v in headers]
                        datasets[sheet_name]=[dict(zip(keys,row)) for row in rows if any(v is not None and str(v).strip() for v in row)]
            wb.close()
        elif filename.endswith('.csv'):
            text_data=uploaded.stream.read().decode('utf-8-sig')
            rows=list(csv.DictReader(io.StringIO(text_data)))
            # CSV imports are treated as student records; use the template headers.
            datasets={'Students':[{str(k or '').strip().lower().replace(' ','_'):v for k,v in row.items()} for row in rows]}
        else:
            flash('Unsupported file type. Upload .xlsx or .csv.', 'error')
            return redirect(url_for('admin_dashboard'))
        con=get_db(); added_students=added_workers=skipped=0; errors=[]
        def val(row,*names):
            for name in names:
                value=row.get(name)
                if value is not None and str(value).strip(): return str(value).strip()
            return ''
        for sheet_name, rows in datasets.items():
            for line_no,row in enumerate(rows, start=2):
                name=val(row,'name','full_name'); email=val(row,'email','campus_email').lower(); password=val(row,'password','temporary_password'); mobile=re.sub(r'\D','',val(row,'mobile','phone'))
                if not name or not valid_campus_email(email) or len(password)<8:
                    skipped+=1
                    if len(errors)<8: errors.append(f'{sheet_name} row {line_no}: name, valid @campus.360 email, and password (8+ characters) are required.')
                    continue
                exists=con.execute('SELECT 1 FROM users WHERE email=?',(email,)).fetchone() if sheet_name=='Students' else con.execute('SELECT 1 FROM workers WHERE email=?',(email,)).fetchone()
                if exists:
                    skipped+=1
                    if len(errors)<8: errors.append(f'{sheet_name} row {line_no}: email already exists; skipped.')
                    continue
                if sheet_name=='Students':
                    next_id=con.execute("SELECT COALESCE(MAX(id),0)+1 n FROM users").fetchone()['n']
                    student_id=val(row,'student_id','studentid') or ('CC'+str(2026000+next_id))
                    con.execute('INSERT INTO users(name,email,password,role,room,mobile,student_id,department,year) VALUES(?,?,?,?,?,?,?,?,?)',(name,email,generate_password_hash(password),'student',val(row,'room','hostel_room'),mobile,val(row,'student_id','studentid') or student_id,val(row,'department','dept'),val(row,'year','academic_year')))
                    added_students+=1
                elif sheet_name=='Workers':
                    department=val(row,'department','dept')
                    if not department or len(mobile)!=10:
                        skipped+=1
                        if len(errors)<8: errors.append(f'{sheet_name} row {line_no}: department and 10-digit mobile are required.')
                        continue
                    code=val(row,'worker_code','worker_id') or ('WK'+uuid.uuid4().hex[:6].upper())
                    con.execute('INSERT INTO workers(name,department,mobile,active,email,password,worker_code) VALUES(?,?,?,?,?,?,?)',(name,department,mobile,1,email,generate_password_hash(password),code))
                    added_workers+=1
                else:
                    skipped+=1
        con.commit(); con.close()
        summary=f'Import complete: {added_students} student(s), {added_workers} worker(s) added; {skipped} row(s) skipped.'
        flash(summary,'success' if not skipped else 'error')
        for issue in errors:
            flash(issue,'error')
    except Exception as exc:
        logging.exception('Admin Excel import failed')
        try: con.rollback(); con.close()
        except Exception: pass
        flash('Import failed. Check the workbook sheet names and column headers, then try again.','error')
    return redirect(url_for('admin_dashboard'))


@app.route('/admin/students', methods=['GET','POST'])
def admin_students():
    if not require('admin'): return redirect(url_for('login'))
    con=get_db()
    if request.method == 'POST':
        name=request.form.get('name','').strip(); email=request.form.get('email','').strip().lower(); password=request.form.get('password',''); mobile=re.sub(r'\D','',request.form.get('mobile','')); room=request.form.get('room','').strip(); dept=request.form.get('department','').strip(); year=request.form.get('year','').strip()
        if not name or not valid_campus_email(email) or len(password)<8 or len(mobile)!=10:
            flash('Enter a name, @campus.360 email, password of at least 8 characters, and 10-digit mobile.', 'error')
        else:
            try:
                next_id=con.execute("SELECT COALESCE(MAX(id),0)+1 n FROM users").fetchone()['n']
                sid='CC'+str(2026000+next_id)
                con.execute('INSERT INTO users(name,email,password,role,room,mobile,student_id,department,year) VALUES(?,?,?,?,?,?,?,?,?)',(name,email,generate_password_hash(password),'student',room,mobile,sid,dept,year)); con.commit()
                flash('Student account created. Student ID: '+sid, 'success')
            except Exception:
                con.rollback()
                flash('That email is already registered or the database rejected the request.', 'error')
    q=request.args.get('q','').strip(); sql="SELECT u.*,COUNT(c.id) complaints_count,SUM(CASE WHEN c.status='Resolved' THEN 1 ELSE 0 END) resolved_count FROM users u LEFT JOIN complaints c ON c.user_id=u.id WHERE u.role='student'"; params=[]
    if q: sql+=" AND (u.name LIKE ? OR u.email LIKE ? OR u.mobile LIKE ? OR u.student_id LIKE ? OR u.department LIKE ?)"; params += [f'%{q}%']*5
    sql+=' GROUP BY u.id ORDER BY u.created_at DESC'; students=con.execute(sql,params).fetchall(); con.close()
    body=render_template_string(r'''<div class="page-head"><div><div class="eyebrow">REGISTERED STUDENTS</div><h1>Student directory</h1><p>Create credentials centrally. Students cannot self-register.</p></div></div><div class="card" style="margin-bottom:20px"><h3>Create student login</h3><form method="POST" class="form-grid"><div><label>Full name</label><input name="name" required></div><div><label>Campus email (@campus.360)</label><input type="email" name="email" pattern="[^@]+@campus\.360" required></div><div><label>Temporary password (8+ chars)</label><input type="password" name="password" minlength="8" required></div><div><label>Mobile</label><input name="mobile" pattern="[0-9]{10}" maxlength="10" required></div><div><label>Hostel / Room</label><input name="room"></div><div><label>Department</label><input name="department"></div><div><label>Year</label><select name="year"><option value="">Choose year</option><option>1st Year</option><option>2nd Year</option><option>3rd Year</option><option>4th Year</option></select></div><div style="display:flex;align-items:end"><button class="btn btn-primary" type="submit">Create student credentials →</button></div></form></div><form class="searchbar"><input name="q" value="{{q}}" placeholder="Search name, student ID, email, mobile or department…"><button class="btn btn-blue">Search</button></form><div class="table-wrap"><table><thead><tr><th>Student</th><th>Student ID</th><th>Contact</th><th>Academic</th><th>Hostel</th><th>Complaints</th><th>Registered</th></tr></thead><tbody>{% for s in students %}<tr><td><strong>{{s['name']}}</strong><br><span class="muted">{{s['email']}}</span></td><td>{{s['student_id'] or '—'}}</td><td>{{s['mobile'] or '—'}}</td><td>{{s['department'] or '—'}}<br>{{s['year'] or ''}}</td><td>{{s['room'] or '—'}}</td><td>{{s['complaints_count'] or 0}} total<br>{{s['resolved_count'] or 0}} resolved</td><td>{{s['created_at'] or '—'}}</td></tr>{% else %}<tr><td colspan="7">No registered students found.</td></tr>{% endfor %}</tbody></table></div>''',students=students,q=q)
    return render_page(body,'Students','STUDENT DIRECTORY')

@app.route('/admin/workers',methods=['GET','POST'])
def admin_workers():
    if not require('admin'): return redirect(url_for('login'))
    con=get_db()
    if request.method=='POST':
        name=request.form.get('name','').strip(); dept=request.form.get('department','').strip(); mobile=re.sub(r'\D','',request.form.get('mobile','')); email=request.form.get('email','').strip().lower(); password=request.form.get('password','')
        if not name or not dept or len(mobile)!=10 or not valid_campus_email(email) or len(password)<8:
            flash('Enter name, department, 10-digit mobile, @campus.360 email and password of at least 8 characters.', 'error')
        else:
            try:
                code='WK'+uuid.uuid4().hex[:6].upper()
                con.execute('INSERT INTO workers(name,department,mobile,email,password,worker_code) VALUES(?,?,?,?,?,?)',(name,dept,mobile,email,generate_password_hash(password),code)); con.commit(); flash('Worker login created. Worker ID: '+code, 'success')
            except Exception:
                con.rollback()
                flash('That worker email is already in use or the database rejected the request.', 'error')
    workers=con.execute('SELECT * FROM workers ORDER BY active DESC,name').fetchall(); con.close()
    body=render_template_string(r'''<div class="page-head"><div><div class="eyebrow">WORKFORCE</div><h1>Workers & staff</h1><p>Create staff accounts and issue their private sign-in credentials.</p></div></div><div class="card" style="margin-bottom:20px"><h3>Create worker login</h3><form class="form-grid" method="POST"><div><label>Worker name</label><input name="name" required placeholder="Worker name"></div><div><label>Department</label><input name="department" required placeholder="Electrical / Plumbing / Cleaning"></div><div><label>Mobile number</label><input name="mobile" required pattern="[0-9]{10}" maxlength="10" placeholder="10-digit number"></div><div><label>Work email (@campus.360)</label><input name="email" type="email" required placeholder="worker@campus.360"></div><div><label>Temporary password (8+ chars)</label><input name="password" type="password" minlength="8" required></div><div style="display:flex;align-items:end"><button class="btn btn-primary" type="submit">＋ Create worker credentials</button></div></form><p class="muted">Existing demo workers have no portal login until you create an account for them here.</p></div><div class="worker-grid">{% for w in workers %}<div class="worker-card"><div class="avatar" style="margin-bottom:12px">{{w['name'][0]|upper}}</div><h3>{{w['name']}}</h3><div class="muted">{{w['department']}}</div><div class="muted">{{w['email'] or 'No portal login assigned'}}</div><div class="muted">Worker ID: {{w['worker_code'] or '—'}}</div><a class="phone" href="tel:{{w['mobile']}}">☎ {{w['mobile']}}</a></div>{% endfor %}</div>''',workers=workers)
    return render_page(body,'Workers','WORKFORCE')

@app.route('/emergency')
def emergency_help():
    body = render_template_string(r'''<div class="page-head"><div><div class="eyebrow">CAMPUS SAFETY</div><h1>Emergency & SOS</h1><p>Quick access to emergency services and campus support contacts.</p></div><a class="btn btn-light" href="{{url_for('home')}}">← Back to dashboard</a></div><section class="emergency-panel"><div class="hero-kicker">IMMEDIATE DANGER</div><h2>Get help now.</h2><p>If someone is in immediate danger, move to a safe place if possible and call the appropriate emergency service. The SOS button opens your phone dialler; it does not silently send an alert.</p><div class="emergency-number">112</div><a class="emergency-call" href="tel:112">🆘 Call Emergency Services · 112</a></section><div class="contact-grid"><div class="contact-card"><div class="tile-icon">🚑</div><h3>Ambulance / Medical</h3><p class="muted">For serious injury or a medical emergency, call 112 and clearly state your location.</p><a class="phone" href="tel:112">☎ Call 112</a></div><div class="contact-card"><div class="tile-icon">🔥</div><h3>Fire Emergency</h3><p class="muted">For fire or smoke, evacuate using safe exits. Do not use lifts or re-enter the building.</p><a class="phone" href="tel:112">☎ Call 112</a></div><div class="contact-card"><div class="tile-icon">🛡️</div><h3>Police / Security</h3><p class="muted">For immediate threats, get to a secure location and contact emergency services or campus security.</p><a class="phone" href="tel:112">☎ Call 112</a></div></div><div class="card" style="margin-top:20px"><h3 style="margin-top:0">Campus support directory</h3><p class="muted">Campus numbers shown here may be demo numbers. Confirm and replace them with your institution’s official contacts before publishing.</p><div class="contact-grid">{% for name,dept,mobile in contacts %}<div class="contact-card"><h3>{{name}}</h3><div class="muted">{{dept}}</div><a class="phone" href="tel:{{mobile}}">☎ {{mobile}}</a></div>{% endfor %}</div></div>''', contacts=CONTACTS)
    return render_page(body, 'Emergency & SOS', 'CAMPUS SAFETY')

@app.route('/contacts')
def contacts():
    body=render_template_string(r'''<div class="page-head"><div><div class="eyebrow">CAMPUS CONTACTS</div><h1>Who can help?</h1><p>Keep important campus support numbers in one place.</p></div></div><div class="contact-grid">{% for name,dept,mobile in contacts %}<div class="contact-card"><div class="avatar">☎</div><h3 style="margin-top:14px">{{name}}</h3><div class="muted">{{dept}}</div><a class="phone" href="tel:{{mobile}}">{{mobile}}</a></div>{% endfor %}</div><div class="card" style="margin-top:20px"><h3 style="margin-top:0">Note</h3><p class="muted">These are demo contact numbers for the project. Replace them with your college's actual contacts before your final presentation.</p></div>''',contacts=CONTACTS)
    return render_page(body,'Contacts','CAMPUS CONTACTS')

@app.route('/help')
def help_support():
    body=render_template_string(r'''<div class="page-head"><div><div class="eyebrow">CAMPUS CARE CONCIERGE</div><h1>Help, without the hassle.</h1><p>Find a quick fix, reach the right team, or ask the Campus Care 360 assistant.</p></div></div>
    <section class="help-hero"><div class="hero-kicker">SUPPORT CENTRE · HERE FOR YOU</div><h2>Let’s get things sorted.</h2><p>Whether you’re reporting a campus issue, checking a complaint, or having trouble signing in, start with one of these quick routes. For immediate danger, use Emergency / SOS from the app menu.</p></section>
    <div class="help-quicklinks"><a class="help-quicklink" href="{{url_for('ai_assistant')}}"><span class="help-quick-icon">✦</span><span><b>Ask Campus Care 360 AI</b><small>Get guided help in seconds</small></span></a><a class="help-quicklink" href="{{url_for('contacts')}}"><span class="help-quick-icon">☎</span><span><b>Contact the right team</b><small>Campus, hostel and support contacts</small></span></a><a class="help-quicklink" href="{{url_for('emergency_help')}}"><span class="help-quick-icon">＋</span><span><b>Emergency help</b><small>Urgent safety guidance and numbers</small></span></a></div>
    <div class="section-title"><h3>Popular fixes</h3><span>Try these first</span></div><div class="help-grid">
    <div class="help-card"><div class="help-card-icon">🔐</div><h3>Can’t sign in?</h3><p class="muted">Check that you’re using the exact email and password issued by the administrator. Use Forgot password to request a reset; an admin must approve it.</p></div>
    <div class="help-card"><div class="help-card-icon">📝</div><h3>Can’t submit a complaint?</h3><p class="muted">Complete the title, category, location and description. If you attach a photo, use PNG, JPG, JPEG, GIF or WEBP and try a smaller image if upload fails.</p></div>
    <div class="help-card"><div class="help-card-icon">🧭</div><h3>Can’t find a complaint?</h3><p class="muted">Open your dashboard and select the complaint card. Keep your complaint ID handy when contacting campus support.</p></div>
    <div class="help-card"><div class="help-card-icon">🛠️</div><h3>Status hasn’t changed?</h3><p class="muted">Refresh the complaint page. If it remains unchanged, contact the administrator and share your complaint ID and latest update.</p></div>
    <div class="help-card"><div class="help-card-icon">📱</div><h3>Page looks broken?</h3><p class="muted">Refresh the page, check your internet connection, and try an updated Chrome or Edge browser. Sign out and sign back in if needed.</p></div>
    <div class="help-card"><div class="help-card-icon">🧑‍💻</div><h3>Still need a hand?</h3><p class="muted">Open the AI Assistant for guided troubleshooting, or use Contacts to reach your campus team. Include a screenshot when reporting technical errors.</p></div></div>
    <div class="section-title" style="margin-top:28px"><h3>Frequently asked questions</h3><span>Helpful details</span></div><div class="faq"><details><summary>Who can create student and worker accounts?</summary><p>Only an authorized administrator should issue account emails, passwords and IDs. If you need access, contact the administrator rather than creating another account.</p></details><details><summary>Where is my complaint information?</summary><p>Sign in and open your dashboard to see your complaints and their latest status. Keep your complaint ID for any follow-up.</p></details><details><summary>What if the app is down?</summary><p>Try again after checking your connection. If the issue persists, contact the administrator and include the page URL, what you clicked and any visible error message.</p></details><details><summary>What should I do in an emergency?</summary><p>If someone is in immediate danger, move to safety if possible and call 112 in India. The app’s SOS shortcut opens the phone dialler; it does not silently send an alert.</p></details></div>''')
    return render_page(body,'Help & Support','SUPPORT CENTRE')

@app.route('/ai', methods=['GET', 'POST'])
def ai_assistant():
    if not session.get('user_id') and not session.get('worker_id'):
        return redirect(url_for('login'))
    account_type = 'worker' if session.get('role') == 'worker' else 'user'
    account_id = session.get('worker_id') if account_type == 'worker' else session.get('user_id')
    con = get_db()
    if request.method == 'POST':
        action = request.form.get('action', 'send')
        if action == 'clear':
            con.execute('DELETE FROM ai_chat_messages WHERE account_type=? AND account_id=?', (account_type, account_id))
            con.commit(); con.close()
            return redirect(url_for('ai_assistant'))
        question = request.form.get('message', '').strip()[:1500]
        if question:
            answer = ai_reply(question)
            con.execute('INSERT INTO ai_chat_messages(account_type,account_id,role,message) VALUES(?,?,?,?)', (account_type, account_id, 'user', question))
            con.execute('INSERT INTO ai_chat_messages(account_type,account_id,role,message) VALUES(?,?,?,?)', (account_type, account_id, 'assistant', answer))
            con.commit()
        con.close()
        return redirect(url_for('ai_assistant'))
    rows = con.execute('SELECT role,message FROM ai_chat_messages WHERE account_type=? AND account_id=? ORDER BY id DESC LIMIT 24', (account_type, account_id)).fetchall()
    con.close()
    history = [{'role': r['role'], 'text': r['message']} for r in reversed(rows)]
    body = render_template_string(r'''<div class="chatgpt-shell">
      <header class="chatgpt-header"><div class="chatgpt-title-wrap"><div class="chatgpt-mark">✦</div><div><h1>Campus Care 360 AI</h1><p>Campus-aware support assistant · Chat history saved to your account</p></div></div><form method="POST"><button class="chatgpt-clear" type="submit" name="action" value="clear">＋ <span>New chat</span></button></form></header>
      <main class="chatgpt-conversation" id="chatgptConversation">
        {% if history %}{% for item in history %}<div class="chatgpt-message {{ 'from-user' if item.role == 'user' else 'from-ai' }}">{% if item.role == 'assistant' %}<div class="chatgpt-avatar">✦</div>{% endif %}<div class="chatgpt-message-content">{% if item.role == 'assistant' %}<div class="chatgpt-speaker">Campus Care 360</div>{% endif %}<div class="chatgpt-message-text">{{ item.text }}</div></div></div>{% endfor %}
        {% else %}<section class="chatgpt-welcome"><div class="welcome-orb">✦</div><h2>How can I help you today?</h2><p>Ask about your complaints, campus notices, lost items, locations or account access.</p><div class="chatgpt-suggestions"><button type="button" onclick="useSuggestion('Show my recent complaints')"><span>⌁</span><b>My complaints</b><small>Check recent status updates</small></button><button type="button" onclick="useSuggestion('What are the latest campus notices?')"><span>▤</span><b>Campus notices</b><small>See recent announcements</small></button><button type="button" onclick="useSuggestion('How do I report a lost item?')"><span>⌕</span><b>Lost & Found</b><small>Report or claim an item</small></button><button type="button" onclick="useSuggestion('How can I reset my password?')"><span>🔐</span><b>Account help</b><small>Request an admin-approved reset</small></button></div></section>{% endif %}
      </main><div class="chatgpt-composer-wrap"><form class="chatgpt-composer" method="POST" id="chatgptForm"><textarea id="chatgptInput" name="message" rows="1" maxlength="1500" required placeholder="Message Campus Care 360…" aria-label="Message Campus Care 360"></textarea><button class="chatgpt-send" type="submit" aria-label="Send message">↑</button></form><p class="chatgpt-disclaimer">Answers use available campus data and may be incomplete. For immediate danger in India, call <a href="tel:112">112</a>.</p></div></div>
    <script>const chatInput=document.getElementById('chatgptInput');function useSuggestion(text){chatInput.value=text;chatInput.focus();}chatInput.addEventListener('input',function(){this.style.height='auto';this.style.height=Math.min(this.scrollHeight,150)+'px';});chatInput.addEventListener('keydown',function(e){if(e.key==='Enter'&&!e.shiftKey){e.preventDefault();document.getElementById('chatgptForm').requestSubmit();}});const conversation=document.getElementById('chatgptConversation');conversation.scrollTop=conversation.scrollHeight;</script>''', history=history)
    return render_page(body, 'AI Assistant', 'AI SUPPORT')

@app.route('/profile')
def student_profile():
    if not require('student'): return redirect(url_for('login'))
    con=get_db(); u=con.execute('SELECT * FROM users WHERE id=?',(session['user_id'],)).fetchone(); count=con.execute('SELECT COUNT(*) n FROM complaints WHERE user_id=?',(session['user_id'],)).fetchone()['n']; resolved=con.execute("SELECT COUNT(*) n FROM complaints WHERE user_id=? AND status IN ('Resolved','Closed')",(session['user_id'],)).fetchone()['n']; active=con.execute("SELECT COUNT(*) n FROM complaints WHERE user_id=? AND status IN ('Submitted','Assigned','In Progress')",(session['user_id'],)).fetchone()['n']; con.close()
    body=render_template_string(r'''<div class="page-head"><div><div class="eyebrow">YOUR CAMPUS IDENTITY</div><h1>My profile</h1><p>Your details, campus identity and support activity in one place.</p></div><a class="btn btn-primary" href="{{url_for('student_dashboard')}}">← Dashboard</a></div>
    <section class="profile-hero"><div class="profile-avatar">{{u['name'][0]|upper}}</div><div><div class="hero-kicker">CAMPUS CARE MEMBER</div><h1>{{u['name']}}</h1><p>{{u['email']}}</p><span class="profile-role">✦ VERIFIED STUDENT PROFILE</span></div></section>
    <div class="profile-grid"><div class="card profile-info-card"><h3>Personal & academic details</h3><div class="profile-details"><div class="profile-detail"><span>Student ID</span><strong>{{u['student_id'] or 'Not assigned'}}</strong></div><div class="profile-detail"><span>Mobile number</span><strong>{{u['mobile'] or 'Not provided'}}</strong></div><div class="profile-detail"><span>Department</span><strong>{{u['department'] or 'Not provided'}}</strong></div><div class="profile-detail"><span>Academic year</span><strong>{{u['year'] or 'Not provided'}}</strong></div><div class="profile-detail"><span>Hostel / room</span><strong>{{u['room'] or 'Not provided'}}</strong></div><div class="profile-detail"><span>Account role</span><strong>Student</strong></div></div><p class="muted" style="margin:17px 0 0">Need to correct your registered details? Contact the campus administrator.</p></div>
    <div><div class="section-title"><h3>Your support activity</h3><span>Live overview</span></div><div class="profile-grid" style="grid-template-columns:1fr 1fr;gap:10px"><div class="profile-stat-card"><small>Total complaints</small><strong>{{count}}</strong><small>Reports submitted</small></div><div class="profile-stat-card"><small>Resolved</small><strong>{{resolved}}</strong><small>Completed cases</small></div><div class="profile-stat-card" style="grid-column:1/-1"><small>In progress</small><strong>{{active}}</strong><small>Reports being handled by campus teams</small></div></div><div class="profile-shortcuts"><a class="profile-shortcut" href="{{url_for('new_complaint')}}"><span>＋</span>Report a new issue <em>›</em></a><a class="profile-shortcut" href="{{url_for('ai_assistant')}}"><span>✦</span>Ask Campus Care 360 AI <em>›</em></a><a class="profile-shortcut" href="{{url_for('help_support')}}"><span>?</span>Help & support <em>›</em></a></div></div></div>''',u=u,count=count,resolved=resolved,active=active)
    return render_page(body,'My Profile','MY PROFILE')

@app.route('/admin/profile')
def admin_profile():
    if not require('admin'): return redirect(url_for('login'))
    con=get_db(); u=con.execute('SELECT * FROM users WHERE id=?',(session['user_id'],)).fetchone(); total=con.execute('SELECT COUNT(*) n FROM complaints').fetchone()['n']; resolved=con.execute("SELECT COUNT(*) n FROM complaints WHERE status IN ('Resolved','Closed')").fetchone()['n']; students=con.execute("SELECT COUNT(*) n FROM users WHERE role='student'").fetchone()['n']; workers=con.execute('SELECT COUNT(*) n FROM workers WHERE active=1').fetchone()['n']; con.close()
    body=render_template_string(r'''<div class="page-head"><div><div class="eyebrow">ADMINISTRATOR IDENTITY</div><h1>Admin profile</h1><p>Account information and a quick view of campus support operations.</p></div><span class="chip" style="background:#f5e4df;color:#8f202a;padding:9px 12px">● ADMIN ACCESS</span></div>
    <section class="profile-hero"><div class="profile-avatar">{{u['name'][0]|upper}}</div><div><div class="hero-kicker">CAMPUS CARE · CONTROL CENTRE</div><h1>{{u['name']}}</h1><p>{{u['email']}} · {{u['department'] or 'Campus Administration'}}</p><span class="profile-role">✦ SYSTEM ADMINISTRATOR</span></div></section>
    <div class="profile-grid"><div class="card profile-info-card"><h3>Administrator details</h3><div class="profile-details"><div class="profile-detail"><span>Full name</span><strong>{{u['name']}}</strong></div><div class="profile-detail"><span>Email address</span><strong>{{u['email']}}</strong></div><div class="profile-detail"><span>Mobile</span><strong>{{u['mobile'] or 'Not provided'}}</strong></div><div class="profile-detail"><span>Department</span><strong>{{u['department'] or 'Campus Administration'}}</strong></div><div class="profile-detail"><span>Access level</span><strong>Administrator</strong></div><div class="profile-detail"><span>Account status</span><strong>Active session</strong></div></div><p class="muted" style="margin:17px 0 0">Protect administrator credentials. Only issue student and worker access to verified campus members.</p></div>
    <div><div class="section-title"><h3>Operations snapshot</h3><span>Current totals</span></div><div class="profile-grid" style="grid-template-columns:1fr 1fr;gap:10px"><div class="profile-stat-card"><small>Total complaints</small><strong>{{total}}</strong><small>All submitted reports</small></div><div class="profile-stat-card"><small>Resolved / closed</small><strong>{{resolved}}</strong><small>Completed reports</small></div><div class="profile-stat-card"><small>Students</small><strong>{{students}}</strong><small>Student accounts</small></div><div class="profile-stat-card"><small>Active workers</small><strong>{{workers}}</strong><small>Campus workforce</small></div></div><div class="profile-shortcuts"><a class="profile-shortcut" href="{{url_for('admin_dashboard')}}"><span>▦</span>Manage complaint queue <em>›</em></a><a class="profile-shortcut" href="{{url_for('admin_students')}}"><span>♙</span>Student directory <em>›</em></a><a class="profile-shortcut" href="{{url_for('admin_workers')}}"><span>⚒</span>Workers & staff <em>›</em></a><a class="profile-shortcut" href="{{url_for('admin_password_resets')}}"><span>🔐</span>Password reset requests <em>›</em></a></div></div></div>''',u=u,total=total,resolved=resolved,students=students,workers=workers)
    return render_page(body,'Admin Profile','ADMIN PROFILE')

@app.route('/uploads/<path:filename>')
def uploaded_file(filename): return send_from_directory(UPLOAD_FOLDER,filename)


# ---------------- Student Life: notices, notifications, lost & found, campus map ----------------
def create_notification(account_type, account_id, title, message, link=''):
    'Store a notification for a user or worker account.'
    if not account_id:
        return
    con = get_db()
    con.execute('INSERT INTO notifications(account_type,account_id,title,message,link) VALUES(?,?,?,?,?)',
                (account_type, int(account_id), str(title)[:160], str(message)[:700], str(link)[:500]))
    con.commit(); con.close()


def account_identity():
    if session.get('role') == 'worker':
        return 'worker', session.get('worker_id')
    return 'user', session.get('user_id')


@app.route('/notices', methods=['GET', 'POST'])
def notices_page():
    if not session.get('user_id') and not session.get('worker_id'):
        return redirect(url_for('login'))
    if request.method == 'POST':
        if not require('admin'):
            flash('Only administrators can publish campus notices.', 'error')
            return redirect(url_for('notices_page'))
        title = request.form.get('title', '').strip()[:160]
        body_text = request.form.get('body', '').strip()[:3000]
        category = request.form.get('category', 'General').strip()[:60]
        priority = request.form.get('priority', 'Normal')
        if priority not in ('Normal', 'Important', 'Urgent'):
            priority = 'Normal'
        if not title or not body_text:
            flash('Add a notice title and message.', 'error')
            return redirect(url_for('notices_page'))
        con = get_db()
        con.execute('INSERT INTO notices(title,body,category,priority,created_by) VALUES(?,?,?,?,?)',
                    (title, body_text, category or 'General', priority, session.get('user_id')))
        users = con.execute("SELECT id FROM users WHERE role='student'").fetchall()
        con.commit(); con.close()
        for u in users:
            create_notification('user', u['id'], 'New campus notice', title, url_for('notices_page'))
        flash('Campus notice published and students notified.', 'success')
        return redirect(url_for('notices_page'))
    con = get_db()
    notices = con.execute('SELECT * FROM notices WHERE active=1 ORDER BY created_at DESC LIMIT 100').fetchall()
    con.close()
    body = render_template_string(r'''<div class="page-head"><div><div class="eyebrow">CAMPUS BULLETIN</div><h1>Smart campus notices</h1><p>Official announcements, events and urgent campus updates in one place.</p></div></div>
    {% if session.get('role')=='admin' %}<div class="card" style="margin-bottom:22px"><h3 style="margin-top:0">Publish an announcement</h3><form method="POST" class="form-grid"><div><label>Notice title</label><input name="title" required maxlength="160" placeholder="Example: Library hours updated"></div><div><label>Category</label><select name="category"><option>General</option><option>Academic</option><option>Hostel</option><option>Mess</option><option>Events</option><option>Safety</option><option>Transport</option></select></div><div><label>Priority</label><select name="priority"><option>Normal</option><option>Important</option><option>Urgent</option></select></div><div class="full"><label>Message</label><textarea name="body" required maxlength="3000" placeholder="Write clear, verified information for students"></textarea></div><div><button class="btn btn-primary" type="submit">Publish notice →</button></div></form></div>{% endif %}
    <div class="admin-list">{% for n in notices %}<article class="admin-card"><div class="admin-top"><div><div class="admin-meta">{{n['category']}} · {{n['created_at']}}</div><div class="admin-title">{{n['title']}}</div></div><span class="chip {{'emergency' if n['priority']=='Urgent' else ('progress' if n['priority']=='Important' else 'submitted')}}">{{n['priority']}}</span></div><p class="admin-desc" style="white-space:pre-wrap">{{n['body']}}</p></article>{% else %}<div class="empty">No notices have been published yet.</div>{% endfor %}</div>''', notices=notices)
    return render_page(body, 'Campus Notices', 'CAMPUS BULLETIN')


@app.route('/notifications', methods=['GET', 'POST'])
def notifications_page():
    if not session.get('user_id') and not session.get('worker_id'):
        return redirect(url_for('login'))
    kind, ident = account_identity()
    con = get_db()
    if request.method == 'POST' and request.form.get('action') == 'read_all':
        con.execute('UPDATE notifications SET is_read=1 WHERE account_type=? AND account_id=?', (kind, ident))
        con.commit()
    items = con.execute('SELECT * FROM notifications WHERE account_type=? AND account_id=? ORDER BY created_at DESC LIMIT 100', (kind, ident)).fetchall()
    con.close()
    body = render_template_string(r'''<div class="page-head"><div><div class="eyebrow">YOUR UPDATES</div><h1>Notifications</h1><p>Complaint progress, new notices and important account updates.</p></div><form method="POST"><input type="hidden" name="action" value="read_all"><button class="btn btn-light" type="submit">Mark all read</button></form></div><div class="admin-list">{% for n in items %}<article class="admin-card" style="border-left:4px solid {{'#A51F2B' if not n['is_read'] else '#ded4ca'}}"><div class="admin-top"><div><div class="admin-title">{{n['title']}}</div><p class="admin-desc">{{n['message']}}</p><div class="admin-meta">{{n['created_at']}}</div></div>{% if not n['is_read'] %}<span class="chip emergency">NEW</span>{% endif %}</div>{% if n['link'] %}<a class="btn btn-light" style="margin-top:12px" href="{{n['link']}}">Open update →</a>{% endif %}</article>{% else %}<div class="empty">You're all caught up. New updates will appear here.</div>{% endfor %}</div>''', items=items)
    if request.method == 'POST':
        flash('Notifications marked as read.', 'success')
    return render_page(body, 'Notifications', 'YOUR UPDATES')


LOST_FOUND_CATEGORIES = ["Electronics", "ID / Cards", "Books & Notes", "Keys", "Clothing", "Bags", "Jewellery", "Sports", "Other"]


def save_lost_found_photo(photo):
    """Validate a small image and return base64 data + trusted MIME type.

    Storing Lost & Found photos in the database makes them survive Vercel's
    ephemeral /tmp filesystem. Existing image_path files remain supported for
    older listings.
    """
    if not photo or not photo.filename:
        return "", ""
    safe_name = secure_filename(photo.filename) or "item-photo"
    ext = safe_name.rsplit(".", 1)[-1].lower() if "." in safe_name else ""
    mime_by_ext = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "gif": "image/gif", "webp": "image/webp"}
    if ext not in mime_by_ext:
        raise ValueError("Photo must be PNG, JPG, JPEG, GIF or WEBP.")
    raw = photo.read(5 * 1024 * 1024 + 1)
    if len(raw) > 5 * 1024 * 1024:
        raise ValueError("Please choose an image smaller than 5 MB.")
    if not raw:
        raise ValueError("The selected image is empty. Please choose another photo.")
    return base64.b64encode(raw).decode("ascii"), mime_by_ext[ext]


def remove_lost_found_photo(filename):
    if not filename:
        return
    # Only remove files with the generated Lost & Found prefix.
    if not str(filename).startswith("lf_"):
        return
    try:
        os.remove(os.path.join(UPLOAD_FOLDER, os.path.basename(filename)))
    except OSError:
        pass


@app.route('/lost-found', methods=['GET', 'POST'])
def lost_found_page():
    if not session.get('user_id') and not session.get('worker_id'):
        return redirect(url_for('login'))

    if request.method == 'POST':
        if not require('student'):
            flash('Only student accounts can post or claim a lost/found item.', 'error')
            return redirect(url_for('lost_found_page'))

        action = request.form.get('action', 'create')
        con = get_db()
        if action == 'claim':
            item_id = request.form.get('item_id', '')
            item = con.execute("SELECT * FROM lost_found WHERE id=? AND status='Open'", (item_id,)).fetchone() if item_id.isdigit() else None
            if item and item['user_id'] != session.get('user_id'):
                owner_id, item_title = item['user_id'], item['title']
                con.close()
                create_notification('user', owner_id, 'Lost & Found claim request',
                                    f"{session.get('name','A student')} is asking about your listing: {item_title}.",
                                    url_for('lost_found_page', mine='1'))
                flash('Claim request sent to the poster. Their contact details remain private.', 'success')
            else:
                con.close()
                flash('That item is no longer available or belongs to your account.', 'error')
            return redirect(url_for('lost_found_page'))

        if action == 'create':
            item_type = request.form.get('item_type', 'Lost')
            category = request.form.get('category', 'Other')
            title = request.form.get('title', '').strip()[:140]
            desc = request.form.get('description', '').strip()[:1800]
            hint = request.form.get('description_hint', '').strip()[:120]
            if hint:
                desc = (desc + ('\n\nIdentifying detail: ' if desc else 'Identifying detail: ') + hint)[:2000]
            location = request.form.get('location', '').strip()[:200]
            item_date = request.form.get('item_date', '').strip()[:20]
            if item_type not in ('Lost', 'Found') or category not in LOST_FOUND_CATEGORIES or not title or not desc or not location:
                con.close()
                flash('Complete the item type, category, name, description and location.', 'error')
                return redirect(url_for('lost_found_page'))
            try:
                image_data, image_mime = save_lost_found_photo(request.files.get('item_photo'))
            except ValueError as exc:
                con.close(); flash(str(exc), 'error')
                return redirect(url_for('lost_found_page'))
            con.execute('INSERT INTO lost_found(user_id,item_type,title,description,location,item_date,category,image_path,image_data,image_mime) VALUES(?,?,?,?,?,?,?,?,?,?)',
                        (session['user_id'], item_type, title, desc, location, item_date, category, '', image_data, image_mime))
            con.commit(); con.close()
            flash('Your listing has been posted successfully.', 'success')
            return redirect(url_for('lost_found_page', mine='1'))
        con.close()
        flash('That action is not available.', 'error')
        return redirect(url_for('lost_found_page'))

    query = request.args.get('q', '').strip()[:100]
    item_type_filter = request.args.get('type', '').strip()
    category_filter = request.args.get('category', '').strip()
    mine = request.args.get('mine') == '1' and require('student')
    con = get_db()
    sql = "SELECT lf.*,u.name poster_name FROM lost_found lf JOIN users u ON u.id=lf.user_id WHERE 1=1"
    params = []
    if mine:
        sql += ' AND lf.user_id=?'
        params.append(session['user_id'])
    else:
        sql += " AND lf.status='Open'"
    if query:
        sql += ' AND (lf.title LIKE ? OR lf.description LIKE ? OR lf.location LIKE ?)'
        params += [f'%{query}%'] * 3
    if item_type_filter in ('Lost', 'Found'):
        sql += ' AND lf.item_type=?'; params.append(item_type_filter)
    if category_filter in LOST_FOUND_CATEGORIES:
        sql += ' AND lf.category=?'; params.append(category_filter)
    sql += ' ORDER BY CASE WHEN lf.status=\'Open\' THEN 0 ELSE 1 END, lf.created_at DESC LIMIT 100'
    items = con.execute(sql, params).fetchall()
    if require('student'):
        mine_count = con.execute('SELECT COUNT(*) AS n FROM lost_found WHERE user_id=?', (session['user_id'],)).fetchone()['n']
    else:
        mine_count = 0
    open_count = con.execute("SELECT COUNT(*) AS n FROM lost_found WHERE status='Open'").fetchone()['n']
    con.close()
    body = render_template_string(r'''<style>
    .lf-stats{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:12px;margin-bottom:22px}
    .lf-stat{padding:18px 20px;border:1px solid #e8dcd2;background:#fffdf9;border-radius:18px;box-shadow:0 8px 24px rgba(42,18,18,.035)}
    .lf-stat strong{display:block;font-size:27px;color:#941b2b;letter-spacing:-.7px}.lf-stat span{display:block;font-size:11px;color:#81756d;margin-top:4px}
    .lf-switch{display:flex;gap:8px;align-items:center;flex-wrap:wrap;margin:0 0 16px}.lf-switch a{padding:10px 14px;border:1px solid #e7dcd2;border-radius:12px;background:#fffdf9;font-size:12px;font-weight:800}.lf-switch a.active{background:#941b2b;color:white;border-color:#941b2b}
    .lf-filters{display:grid;grid-template-columns:minmax(180px,2fr) 1fr 1fr auto;gap:10px;margin-bottom:18px}
    .lf-list{display:grid;gap:11px}.lf-row{display:grid;grid-template-columns:82px minmax(0,1fr) auto;gap:16px;align-items:center;background:#fffdf9;border:1px solid #e8dcd2;border-radius:18px;padding:15px 17px;box-shadow:0 6px 20px rgba(42,18,18,.025)}
    .lf-thumb{width:82px;height:82px;object-fit:cover;border-radius:14px;border:1px solid #eee0d5;background:#f8eee6}.lf-thumb-empty{display:grid;place-items:center;color:#941b2b;font-size:27px;background:linear-gradient(135deg,#fff5ed,#f1dcd1)}
    .lf-row-main{min-width:0}.lf-badges{display:flex;gap:6px;flex-wrap:wrap;margin-bottom:7px}.lf-row-title{font-size:14px;font-weight:850;color:#332522}.lf-row-desc{font-size:12px;line-height:1.55;color:#756b66;margin-top:5px;white-space:pre-wrap;overflow-wrap:anywhere}.lf-row-meta{font-size:10px;color:#92857c;margin-top:7px}.lf-row-actions{display:flex;gap:8px;flex-wrap:wrap;justify-content:flex-end;align-items:center}.lf-row-actions form{margin:0}.lf-upload-box{border:1.5px dashed #d7bdb0;border-radius:15px;padding:16px;background:#fffaf5}.lf-upload-box input[type=file]{background:transparent;border:0;padding:4px 0}.lf-preview{display:none;max-width:240px;max-height:190px;object-fit:contain;border-radius:12px;margin-top:12px;border:1px solid #e7dcd2}.lf-note{font-size:11px;color:#81756d;line-height:1.6;margin:8px 0 0}
    @media(max-width:760px){.lf-filters{grid-template-columns:1fr}.lf-stats{grid-template-columns:1fr 1fr}.lf-stat:last-child{grid-column:1/-1}.lf-row{grid-template-columns:62px minmax(0,1fr);gap:12px;padding:12px}.lf-thumb{width:62px;height:62px}.lf-row-actions{grid-column:1/-1;justify-content:flex-start}.lf-row-title{font-size:13px}}
    </style>
    <div class="page-head"><div><div class="eyebrow">COMMUNITY CARE · STUDENT EXCHANGE</div><h1>Lost & Found</h1><p>Use the same simple report form as campus complaints to report a lost or found item, add a photo and follow up on listings.</p></div></div>
    <div class="lf-stats"><div class="lf-stat"><strong>{{open_count}}</strong><span>Open items on campus</span></div><div class="lf-stat"><strong>{{items|length}}</strong><span>Items in this view</span></div>{% if session.get('role')=='student' %}<div class="lf-stat"><strong>{{mine_count}}</strong><span>Your total listings</span></div>{% else %}<div class="lf-stat"><strong>Private</strong><span>Student contact details are protected</span></div>{% endif %}</div>
    {% if session.get('role')=='student' %}<div class="card" style="margin-bottom:24px"><div style="margin-bottom:20px"><h3 style="margin:0 0 7px;font-size:17px">＋ Report a lost or found item</h3><p class="muted" style="margin:0;font-size:12px;line-height:1.7">Complete the required fields below. Add a photo if you have one, and avoid sharing passwords or sensitive ID/bank details.</p></div><form method="POST" enctype="multipart/form-data" id="lf-create-form"><input type="hidden" name="action" value="create"><div class="form-grid"><div><label>Report type *</label><select name="item_type" required><option value="Lost">I lost an item</option><option value="Found">I found an item</option></select></div><div><label>Category *</label><select name="category" required><option value="">Choose a category</option>{% for c in categories_lf %}<option>{{c}}</option>{% endfor %}</select></div><div class="full"><label>Item name / title *</label><input name="title" required maxlength="140" placeholder="Example: Black water bottle, notebook, keys"></div><div><label>Location *</label><input name="location" required maxlength="200" placeholder="Library, Block B, canteen…"></div><div><label>Date lost / found</label><input type="date" name="item_date" max="{{today}}"></div><div class="full"><label>Description *</label><textarea name="description" required maxlength="2000" placeholder="Describe the item, its colour or identifying features, and where it was last seen or found."></textarea></div><div class="full"><label>Additional identifying detail (optional)</label><input name="description_hint" maxlength="120" placeholder="Brand, colour, initials or a distinguishing mark"></div><div class="full lf-upload-box"><label>Item photo (optional)</label><input id="lf-photo-input" type="file" name="item_photo" accept="image/png,image/jpeg,image/gif,image/webp"><img id="lf-photo-preview" class="lf-preview" alt="Selected item photo preview"><p class="lf-note">Accepted formats: PNG, JPG, JPEG, GIF or WEBP. Maximum size: 5 MB. Avoid uploading faces, full ID cards or bank details.</p></div></div><div class="form-actions"><a class="btn btn-light" href="{{url_for('student_dashboard')}}">Cancel</a><button class="btn btn-primary" type="submit">Submit item report →</button></div></form></div>{% endif %}
    <div class="section-title"><h3>{{'My listings' if mine else 'Campus lost & found reports'}}</h3><span>{{items|length}} shown · {{open_count}} open</span></div>
    <div class="lf-switch"><a href="{{url_for('lost_found_page')}}" class="{{'active' if not mine else ''}}">Browse open items</a>{% if session.get('role')=='student' %}<a href="{{url_for('lost_found_page',mine='1')}}" class="{{'active' if mine else ''}}">My listings</a>{% endif %}</div>
    <form class="lf-filters" method="GET"><input name="q" value="{{query}}" placeholder="Search item, description or location"><select name="type"><option value="">Lost and found</option><option value="Lost" {{'selected' if type_filter=='Lost' else ''}}>Lost items</option><option value="Found" {{'selected' if type_filter=='Found' else ''}}>Found items</option></select><select name="category"><option value="">All categories</option>{% for c in categories_lf %}<option value="{{c}}" {{'selected' if category_filter==c else ''}}>{{c}}</option>{% endfor %}</select>{% if mine %}<input type="hidden" name="mine" value="1">{% endif %}<button class="btn btn-light" type="submit">Search / Filter</button></form>
    <div class="lf-list">{% for item in items %}<article class="lf-row">{% if item['image_data'] or item['image_path'] %}<img class="lf-thumb" src="{{url_for('lost_found_image',item_id=item['id'])}}" alt="Photo of {{item['title']}}" loading="lazy">{% else %}<div class="lf-thumb lf-thumb-empty">{% if item['category']=='Electronics' %}▣{% elif item['category']=='Books & Notes' %}▤{% elif item['category']=='Keys' %}⚿{% elif item['category']=='Bags' %}▱{% else %}⌕{% endif %}</div>{% endif %}<div class="lf-row-main"><div class="lf-badges"><span class="chip {{'emergency' if item['item_type']=='Lost' else 'resolved'}}">{{item['item_type']}}</span><span class="chip submitted">{{item['category'] or 'Other'}}</span><span class="chip {{'closed' if item['status']!='Open' else 'progress'}}">{{item['status']}}</span></div><div class="lf-row-title">{{item['title']}}</div><div class="lf-row-desc">{{item['description']}}</div><div class="lf-row-meta">📍 {{item['location']}} · Posted by {{item['poster_name']}} · {{item['created_at']}}{% if item['item_date'] %} · Date lost/found: {{item['item_date']}}{% endif %}</div></div><div class="lf-row-actions">{% if session.get('role')=='student' and item['user_id'] != session.get('user_id') and item['status']=='Open' %}<form method="POST"><input type="hidden" name="action" value="claim"><input type="hidden" name="item_id" value="{{item['id']}}"><button class="btn btn-primary" type="submit">Request to claim</button></form>{% endif %}{% if session.get('role')=='student' and item['user_id']==session.get('user_id') and item['status']=='Open' %}<a class="btn btn-light" href="{{url_for('lost_found_edit',item_id=item['id'])}}">Edit</a><form method="POST" action="{{url_for('lost_found_close',item_id=item['id'])}}" onsubmit="return confirm('Mark this item as resolved?')"><button class="btn btn-light" type="submit">Mark resolved</button></form>{% endif %}{% if session.get('role')=='admin' and item['status']=='Open' %}<form method="POST" action="{{url_for('lost_found_close',item_id=item['id'])}}"><button class="btn btn-light" type="submit">Close listing</button></form>{% endif %}</div></article>{% else %}<div class="empty">No listings match these filters. Try another search or publish a new listing.</div>{% endfor %}</div>
    <script>(function(){const input=document.getElementById('lf-photo-input');const preview=document.getElementById('lf-photo-preview');if(!input||!preview)return;input.addEventListener('change',function(){const f=input.files&&input.files[0];if(!f){preview.style.display='none';preview.removeAttribute('src');return;}if(f.size>5*1024*1024){alert('Please choose an image smaller than 5 MB.');input.value='';preview.style.display='none';return;}if(!['image/png','image/jpeg','image/gif','image/webp'].includes(f.type)){alert('Choose a PNG, JPG, GIF or WEBP image.');input.value='';preview.style.display='none';return;}preview.src=URL.createObjectURL(f);preview.style.display='block';});})();</script>''', items=items, query=query, type_filter=item_type_filter, category_filter=category_filter, mine=mine,
       categories_lf=LOST_FOUND_CATEGORIES, mine_count=mine_count, open_count=open_count,
       today=__import__('datetime').date.today().isoformat())
    return render_page(body, 'Lost & Found', 'COMMUNITY CARE')


@app.route('/lost-found/<int:item_id>/photo')
def lost_found_image(item_id):
    """Serve database-backed Lost & Found photos (persistent on Vercel)."""
    if not session.get('user_id') and not session.get('worker_id'):
        abort(401)
    con = get_db()
    item = con.execute('SELECT image_data,image_mime,image_path FROM lost_found WHERE id=?', (item_id,)).fetchone()
    con.close()
    if not item:
        abort(404)
    if item['image_data']:
        try:
            raw = base64.b64decode(item['image_data'], validate=True)
        except (ValueError, TypeError):
            abort(404)
        mime = item['image_mime'] if item['image_mime'] in ('image/png','image/jpeg','image/gif','image/webp') else 'application/octet-stream'
        return Response(raw, mimetype=mime, headers={'X-Content-Type-Options':'nosniff', 'Cache-Control':'private, max-age=3600'})
    if item['image_path']:
        return send_from_directory(UPLOAD_FOLDER, item['image_path'])
    abort(404)


@app.route('/lost-found/<int:item_id>/edit', methods=['GET', 'POST'])
def lost_found_edit(item_id):
    if not require('student'):
        return redirect(url_for('login'))
    con = get_db()
    item = con.execute('SELECT * FROM lost_found WHERE id=? AND user_id=?', (item_id, session['user_id'])).fetchone()
    if not item:
        con.close(); flash('Listing not found or you do not own it.', 'error')
        return redirect(url_for('lost_found_page', mine='1'))
    if item['status'] != 'Open':
        con.close(); flash('Resolved listings cannot be edited.', 'error')
        return redirect(url_for('lost_found_page', mine='1'))
    if request.method == 'POST':
        item_type = request.form.get('item_type', 'Lost')
        category = request.form.get('category', 'Other')
        title = request.form.get('title', '').strip()[:140]
        description = request.form.get('description', '').strip()[:2000]
        location = request.form.get('location', '').strip()[:200]
        item_date = request.form.get('item_date', '').strip()[:20]
        if item_type not in ('Lost', 'Found') or category not in LOST_FOUND_CATEGORIES or not title or not description or not location:
            con.close(); flash('Complete all required listing fields.', 'error')
            return redirect(url_for('lost_found_edit', item_id=item_id))
        image_data = item['image_data'] or ''
        image_mime = item['image_mime'] or ''
        image_path = item['image_path'] or ''
        try:
            uploaded_data, uploaded_mime = save_lost_found_photo(request.files.get('item_photo'))
            if uploaded_data:
                image_data, image_mime, image_path = uploaded_data, uploaded_mime, ''
            elif request.form.get('remove_photo') == 'yes':
                image_data, image_mime, image_path = '', '', ''
        except ValueError as exc:
            con.close(); flash(str(exc), 'error')
            return redirect(url_for('lost_found_edit', item_id=item_id))
        con.execute('UPDATE lost_found SET item_type=?,category=?,title=?,description=?,location=?,item_date=?,image_path=?,image_data=?,image_mime=? WHERE id=? AND user_id=?',
                    (item_type, category, title, description, location, item_date, image_path, image_data, image_mime, item_id, session['user_id']))
        con.commit(); con.close()
        if image_path != (item['image_path'] or ''):
            remove_lost_found_photo(item['image_path'])
        flash('Your listing has been updated.', 'success')
        return redirect(url_for('lost_found_page', mine='1'))
    con.close()
    body = render_template_string(r'''<div class="page-head"><div><div class="eyebrow">YOUR LISTING</div><h1>Edit item</h1><p>Update the details or replace the item photo.</p></div><a class="btn btn-light" href="{{url_for('lost_found_page',mine='1')}}">← My listings</a></div><div class="card"><form method="POST" enctype="multipart/form-data" class="form-grid"><div><label>Listing type</label><select name="item_type"><option {{'selected' if item['item_type']=='Lost' else ''}}>Lost</option><option {{'selected' if item['item_type']=='Found' else ''}}>Found</option></select></div><div><label>Category</label><select name="category">{% for c in categories_lf %}<option {{'selected' if (item['category'] or 'Other')==c else ''}}>{{c}}</option>{% endfor %}</select></div><div><label>Item name</label><input name="title" required maxlength="140" value="{{item['title']}}"></div><div><label>Location</label><input name="location" required maxlength="200" value="{{item['location']}}"></div><div><label>Date lost/found</label><input type="date" name="item_date" value="{{item['item_date'] or ''}}"></div><div class="full"><label>Description</label><textarea name="description" required maxlength="2000">{{item['description']}}</textarea></div>{% if item['image_data'] or item['image_path'] %}<div class="full"><label>Current photo</label><img src="{{url_for('lost_found_image',item_id=item['id'])}}" alt="Current item photo" style="display:block;max-width:260px;max-height:220px;object-fit:contain;border-radius:14px;border:1px solid #e7dcd2"><label style="margin-top:10px"><input type="checkbox" name="remove_photo" value="yes" style="width:auto"> Remove current photo if no replacement is uploaded</label></div>{% endif %}<div class="full"><label>Replace photo (optional, up to 5 MB)</label><input type="file" name="item_photo" accept="image/png,image/jpeg,image/gif,image/webp"></div><div class="full"><button class="btn btn-primary" type="submit">Save changes →</button></div></form></div>''', item=item, categories_lf=LOST_FOUND_CATEGORIES)
    return render_page(body, 'Edit Lost & Found Item', 'COMMUNITY CARE')


@app.route('/lost-found/<int:item_id>/close', methods=['POST'])
def lost_found_close(item_id):
    if not session.get('user_id') and not session.get('worker_id'):
        return redirect(url_for('login'))
    con = get_db()
    if require('admin'):
        item = con.execute("SELECT id,user_id,title FROM lost_found WHERE id=? AND status='Open'", (item_id,)).fetchone()
    elif require('student'):
        item = con.execute("SELECT id,user_id,title FROM lost_found WHERE id=? AND user_id=? AND status='Open'", (item_id, session['user_id'])).fetchone()
    else:
        item = None
    if not item:
        con.close(); flash('Listing not found or you do not have permission to close it.', 'error')
        return redirect(url_for('lost_found_page'))
    con.execute("UPDATE lost_found SET status='Closed' WHERE id=?", (item_id,))
    con.commit(); con.close()
    flash('Listing marked as resolved.', 'success')
    return redirect(url_for('lost_found_page', mine='1') if require('student') else url_for('lost_found_page'))

def clean_map_form():
    name = request.form.get('name', '').strip()[:140]
    category = request.form.get('category', 'General').strip()[:60]
    building = request.form.get('building', '').strip()[:140]
    description = request.form.get('description', '').strip()[:1000]
    hours = request.form.get('hours', '').strip()[:100]
    phone = re.sub(r'[^0-9+ -]', '', request.form.get('phone', ''))[:30]
    map_url = request.form.get('map_url', '').strip()[:500]
    if map_url and not (map_url.startswith('https://') or map_url.startswith('http://')):
        raise ValueError('Map links must start with http:// or https://.')
    if not name:
        raise ValueError('Enter a location name.')
    return (name, category or 'General', building, description, hours, phone, map_url)


@app.route('/campus-map', methods=['GET', 'POST'])
def campus_map():
    if not session.get('user_id') and not session.get('worker_id'):
        return redirect(url_for('login'))
    if request.method == 'POST':
        if not require('admin'):
            flash('Only administrators can add campus locations.', 'error')
            return redirect(url_for('campus_map'))
        try:
            values = clean_map_form()
        except ValueError as exc:
            flash(str(exc), 'error')
            return redirect(url_for('campus_map'))
        con = get_db()
        con.execute('INSERT INTO campus_locations(name,category,building,description,hours,phone,map_url) VALUES(?,?,?,?,?,?,?)', values)
        con.commit(); con.close()
        flash('Campus location added.', 'success')
        return redirect(url_for('campus_map'))
    query = request.args.get('q', '').strip()[:100]
    con = get_db(); sql = 'SELECT * FROM campus_locations WHERE active=1'; params = []
    if query:
        sql += ' AND (name LIKE ? OR category LIKE ? OR building LIKE ? OR description LIKE ?)'; params += [f'%{query}%'] * 4
    sql += ' ORDER BY category,name'; locations = con.execute(sql, params).fetchall(); con.close()
    body = render_template_string(r'''<div class="page-head"><div><div class="eyebrow">FIND YOUR WAY</div><h1>Campus map & services</h1><p>Find campus buildings and support services. Administrators can add, edit and remove outdated locations.</p></div><span class="chip submitted">{{locations|length}} locations</span></div>
    <div class="help-hero"><div class="hero-kicker">CAMPUS NAVIGATION</div><h2>Know where to go.</h2><p>Search by building, service or category. Open a saved map link, or use Google Maps search for the location name. Only add verified campus details.</p></div>
    {% if session.get('role')=='admin' %}<div class="card" style="margin-bottom:20px"><h3 style="margin-top:0">＋ Add a campus location</h3><form method="POST" class="form-grid"><div><label>Place name *</label><input name="name" required placeholder="Central Library"></div><div><label>Category</label><select name="category"><option>Academic</option><option>Library</option><option>Hostel</option><option>Food & Mess</option><option>Medical</option><option>Security</option><option>Administration</option><option>Sports</option><option>Transport</option><option>Other</option></select></div><div><label>Building / block</label><input name="building" placeholder="Block A"></div><div><label>Opening hours</label><input name="hours" placeholder="9 AM – 6 PM"></div><div><label>Contact phone (optional)</label><input name="phone" placeholder="Verified campus number"></div><div><label>Google Maps link (optional)</label><input name="map_url" type="url" placeholder="https://maps.google.com/…"></div><div class="full"><label>Description / directions</label><textarea name="description" placeholder="Floor, landmark or service details"></textarea></div><div><button class="btn btn-primary" type="submit">Save location →</button></div></form></div>{% endif %}
    <form class="searchbar"><input name="q" value="{{query}}" placeholder="Search library, hostel, medical centre, block…"><button class="btn btn-light">Search</button></form><div class="contact-grid">{% for loc in locations %}<article class="contact-card"><div class="tile-icon">⌖</div><span class="chip submitted">{{loc['category']}}</span><h3 style="margin-top:12px">{{loc['name']}}</h3>{% if loc['building'] %}<p class="muted"><b>Building:</b> {{loc['building']}}</p>{% endif %}<p class="muted">{{loc['description']}}</p>{% if loc['hours'] %}<p class="muted">Hours: {{loc['hours']}}</p>{% endif %}{% if loc['phone'] %}<a class="phone" href="tel:{{loc['phone']}}">☎ {{loc['phone']}}</a>{% endif %}<div style="display:flex;gap:8px;flex-wrap:wrap;margin-top:14px"><a class="btn btn-primary" target="_blank" rel="noopener noreferrer" href="{{loc['map_url'] or ('https://www.google.com/maps/search/?api=1&query=' ~ (loc['name'] ~ ' ' ~ loc['building'])|urlencode)}}">Open map / directions ↗</a>{% if session.get('role')=='admin' %}<a class="btn btn-light" href="#edit-location-{{loc['id']}}">Edit details</a>{% endif %}</div>
    {% if session.get('role')=='admin' %}<details id="edit-location-{{loc['id']}}" style="margin-top:14px;border-top:1px solid #e7dcd2;padding-top:14px"><summary style="font-size:12px;font-weight:850;cursor:pointer">Update this location</summary><form method="POST" action="{{url_for('campus_map_update',location_id=loc['id'])}}" class="form-grid" style="margin-top:12px"><div><label>Place name</label><input name="name" required value="{{loc['name']}}"></div><div><label>Category</label><input name="category" required value="{{loc['category']}}"></div><div><label>Building / block</label><input name="building" value="{{loc['building']}}"></div><div><label>Opening hours</label><input name="hours" value="{{loc['hours']}}"></div><div><label>Phone</label><input name="phone" value="{{loc['phone']}}"></div><div><label>Map link</label><input name="map_url" type="url" value="{{loc['map_url']}}" placeholder="https://…"></div><div class="full"><label>Description / directions</label><textarea name="description">{{loc['description']}}</textarea></div><div><button class="btn btn-primary" type="submit">Save changes</button></div></form><form method="POST" action="{{url_for('campus_map_remove',location_id=loc['id'])}}" onsubmit="return confirm('Remove this location from the campus map?')" style="margin-top:10px"><button class="btn btn-light" type="submit">Remove location</button></form></details>{% endif %}</article>{% else %}<div class="empty">No locations have been added yet. Ask an administrator to add verified campus locations.</div>{% endfor %}</div>''', locations=locations, query=query)
    return render_page(body, 'Campus Map', 'FIND YOUR WAY')


@app.route('/campus-map/<int:location_id>/update', methods=['POST'])
def campus_map_update(location_id):
    if not require('admin'):
        return redirect(url_for('login'))
    try:
        values = clean_map_form()
    except ValueError as exc:
        flash(str(exc), 'error')
        return redirect(url_for('campus_map'))
    con = get_db()
    cur = con.execute('UPDATE campus_locations SET name=?,category=?,building=?,description=?,hours=?,phone=?,map_url=? WHERE id=? AND active=1', values + (location_id,))
    con.commit(); con.close()
    flash('Campus location updated.' if getattr(cur, 'rowcount', 1) else 'Location not found or already removed.', 'success')
    return redirect(url_for('campus_map'))


@app.route('/campus-map/<int:location_id>/remove', methods=['POST'])
def campus_map_remove(location_id):
    if not require('admin'):
        return redirect(url_for('login'))
    con = get_db(); con.execute('UPDATE campus_locations SET active=0 WHERE id=?', (location_id,)); con.commit(); con.close()
    flash('Location removed from the active campus map.', 'success')
    return redirect(url_for('campus_map'))


initialize_database()
if os.environ.get("VERCEL") and not DATABASE_URL:
    logging.warning("DATABASE_URL is not configured: SQLite data in /tmp is temporary and may disappear across Vercel instances. Configure a managed PostgreSQL database for persistent records.")
if __name__ == "__main__":
    app.run(debug=False, host="0.0.0.0", port=int(os.environ.get("PORT", 5000)))
